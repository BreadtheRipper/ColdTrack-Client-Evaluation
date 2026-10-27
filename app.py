from __future__ import annotations

import argparse
import difflib
import getpass
import json
import os
import re
import secrets
import sqlite3
import threading
import time
import zipfile
from datetime import date, datetime
from functools import wraps
from pathlib import Path
from typing import Any

from flask import Flask, g, jsonify, render_template, request, send_file, session
from werkzeug.security import check_password_hash, generate_password_hash


ROOT = Path(__file__).resolve().parent
INSTANCE = ROOT / "instance"
DEFAULT_DATABASE = INSTANCE / "sorter.sqlite3"

LOOKUP_FIELDS = {
    "product_line",
    "equipment_type",
    "equipment_brand",
    "model",
    "store_name",
    "area",
    "province",
    "manager",
    "sales_rep",
    "technician",
}
ASSET_FIELDS = {
    "product_line",
    "equipment_type",
    "equipment_brand",
    "model",
    "serial_number",
    "store_name",
    "business_address",
    "area",
    "province",
    "manager",
    "contact",
    "delivery_date",
    "sales_rep",
    "status",
    "notes",
}
FILTER_FIELDS = {
    "product_line": "a.product_line",
    "equipment_type": "a.equipment_type",
    "equipment_brand": "a.equipment_brand",
    "model": "a.model",
    "serial_number": "a.serial_number",
    "store_name": "a.store_name",
    "business_address": "a.business_address",
    "area": "a.area",
    "province": "a.province",
    "manager": "a.manager",
    "contact": "a.contact",
    "sales_rep": "a.sales_rep",
    "status": "a.status",
}
SORT_FIELDS = {
    "product_line": "a.product_line",
    "equipment_type": "a.equipment_type",
    "equipment_brand": "a.equipment_brand",
    "model": "a.model",
    "serial_number": "a.serial_number",
    "store_name": "a.store_name",
    "business_address": "a.business_address",
    "area": "a.area",
    "province": "a.province",
    "manager": "a.manager",
    "contact": "a.contact",
    "delivery_date": "a.delivery_date",
    "sales_rep": "a.sales_rep",
    "status": "a.status",
    "technician": "(SELECT MIN(se.technician) FROM service_events se WHERE se.asset_id = a.id)",
}

SCHEMA = """
PRAGMA foreign_keys = ON;
CREATE TABLE IF NOT EXISTS users (
    id INTEGER PRIMARY KEY,
    username TEXT NOT NULL UNIQUE COLLATE NOCASE,
    display_name TEXT NOT NULL,
    password_hash TEXT NOT NULL,
    role TEXT NOT NULL CHECK (role IN ('admin', 'staff')),
    active INTEGER NOT NULL DEFAULT 1,
    created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS stores (
    id INTEGER PRIMARY KEY,
    store_name TEXT NOT NULL,
    business_address TEXT NOT NULL DEFAULT '',
    area TEXT NOT NULL DEFAULT '',
    province TEXT NOT NULL DEFAULT '',
    manager TEXT NOT NULL DEFAULT '',
    contact TEXT NOT NULL DEFAULT '',
    sales_rep TEXT NOT NULL DEFAULT '',
    notes TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS assets (
    id INTEGER PRIMARY KEY,
    store_id INTEGER REFERENCES stores(id),
    moved_from_asset_id INTEGER REFERENCES assets(id),
    placement_active INTEGER NOT NULL DEFAULT 1,
    product_line TEXT NOT NULL DEFAULT '',
    source_sheet TEXT,
    source_row INTEGER,
    equipment_type TEXT NOT NULL DEFAULT '',
    equipment_kind TEXT NOT NULL DEFAULT 'freezer',
    equipment_brand TEXT NOT NULL DEFAULT '',
    model TEXT NOT NULL DEFAULT '',
    serial_number TEXT NOT NULL DEFAULT '',
    store_name TEXT NOT NULL DEFAULT '',
    business_address TEXT NOT NULL DEFAULT '',
    area TEXT NOT NULL DEFAULT '',
    province TEXT NOT NULL DEFAULT '',
    manager TEXT NOT NULL DEFAULT '',
    contact TEXT NOT NULL DEFAULT '',
    delivery_date TEXT NOT NULL DEFAULT '',
    sales_rep TEXT NOT NULL DEFAULT '',
    status TEXT NOT NULL DEFAULT 'unassigned'
        CHECK (status IN ('deployed', 'in_transit', 'unassigned')),
    notes TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL,
    UNIQUE (source_sheet, source_row)
);
CREATE INDEX IF NOT EXISTS assets_serial_idx ON assets(serial_number);
CREATE INDEX IF NOT EXISTS assets_area_province_idx ON assets(area, province);
CREATE INDEX IF NOT EXISTS assets_brand_model_idx ON assets(equipment_brand, model);
CREATE TABLE IF NOT EXISTS service_events (
    id INTEGER PRIMARY KEY,
    asset_id INTEGER NOT NULL REFERENCES assets(id) ON DELETE CASCADE,
    service_date TEXT NOT NULL,
    technician TEXT NOT NULL,
    repairs TEXT NOT NULL DEFAULT '',
    parts TEXT NOT NULL DEFAULT '',
    notes TEXT NOT NULL DEFAULT '',
    created_by INTEGER REFERENCES users(id),
    created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS service_asset_date_idx ON service_events(asset_id, service_date);
CREATE INDEX IF NOT EXISTS service_technician_idx ON service_events(technician);
CREATE TABLE IF NOT EXISTS location_events (
    id INTEGER PRIMARY KEY,
    asset_id INTEGER NOT NULL REFERENCES assets(id) ON DELETE CASCADE,
    business_address TEXT NOT NULL DEFAULT '',
    area TEXT NOT NULL DEFAULT '',
    province TEXT NOT NULL DEFAULT '',
    effective_date TEXT,
    notes TEXT NOT NULL DEFAULT '',
    created_by INTEGER REFERENCES users(id),
    created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS location_asset_date_idx ON location_events(asset_id, effective_date);
CREATE TABLE IF NOT EXISTS lookup_values (
    id INTEGER PRIMARY KEY,
    field TEXT NOT NULL,
    value TEXT NOT NULL,
    parent_value TEXT NOT NULL DEFAULT '',
    active INTEGER NOT NULL DEFAULT 1,
    source TEXT NOT NULL DEFAULT 'admin',
    created_by INTEGER REFERENCES users(id),
    UNIQUE (field, value)
);
CREATE INDEX IF NOT EXISTS lookup_field_active_idx ON lookup_values(field, active, value);
CREATE TABLE IF NOT EXISTS app_settings (
    key TEXT PRIMARY KEY,
    value TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS audit_log (
    id INTEGER PRIMARY KEY,
    user_id INTEGER REFERENCES users(id),
    actor TEXT NOT NULL,
    action TEXT NOT NULL,
    entity_type TEXT NOT NULL,
    entity_id TEXT NOT NULL,
    before_json TEXT,
    after_json TEXT,
    created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS audit_created_idx ON audit_log(created_at DESC);
CREATE TABLE IF NOT EXISTS import_batches (
    id INTEGER PRIMARY KEY,
    file_name TEXT NOT NULL,
    sheets_json TEXT NOT NULL,
    records_json TEXT NOT NULL,
    created_by INTEGER REFERENCES users(id),
    status TEXT NOT NULL CHECK (status IN ('review', 'awaiting_approval', 'complete', 'rejected')),
    created_at TEXT NOT NULL,
    completed_at TEXT
);
CREATE TABLE IF NOT EXISTS lookup_requests (
    id INTEGER PRIMARY KEY,
    batch_id INTEGER NOT NULL REFERENCES import_batches(id) ON DELETE CASCADE,
    field TEXT NOT NULL,
    value TEXT NOT NULL,
    status TEXT NOT NULL CHECK (status IN ('pending', 'approved', 'rejected')),
    requested_by INTEGER REFERENCES users(id),
    reviewed_by INTEGER REFERENCES users(id),
    created_at TEXT NOT NULL,
    reviewed_at TEXT,
    UNIQUE (batch_id, field, value)
);
CREATE INDEX IF NOT EXISTS lookup_requests_status_idx ON lookup_requests(status, created_at);
"""


def create_database_backup(database_path: str, backup_directory: str | Path) -> Path:
    directory = Path(backup_directory)
    directory.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().astimezone().strftime("%Y%m%d-%H%M%S")
    target = directory / f"coldtrack-{timestamp}-{secrets.token_hex(3)}.sqlite3"
    temporary = target.with_suffix(target.suffix + ".tmp")
    source = sqlite3.connect(database_path, timeout=30)
    destination = None
    try:
        destination = sqlite3.connect(temporary, timeout=30)
        source.backup(destination)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    finally:
        if destination is not None:
            destination.close()
        source.close()
    os.replace(temporary, target)
    backups = sorted(
        directory.glob("coldtrack-*.sqlite3"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    for old_backup in backups[30:]:
        old_backup.unlink()
    return target


def run_scheduled_backup(application: Flask, now: datetime | None = None) -> bool:
    with application.app_context():
        db = application.get_db()
        settings = {
            row["key"]: row["value"]
            for row in db.execute(
                "SELECT key, value FROM app_settings WHERE key IN ('backup_time', 'backup_last_run_date')"
            ).fetchall()
        }
        current_time = now or datetime.now().astimezone()
        today = current_time.date().isoformat()
        if current_time.strftime("%H:%M") < settings.get("backup_time", "02:00") or settings.get(
            "backup_last_run_date"
        ) == today:
            return False
        create_database_backup(
            application.config["DATABASE"],
            application.config["BACKUP_DIR"],
        )
        db.execute(
            """INSERT INTO app_settings (key, value) VALUES ('backup_last_run_date', ?)
               ON CONFLICT(key) DO UPDATE SET value=excluded.value""",
            (today,),
        )
        db.commit()
        return True


def start_backup_scheduler(application: Flask) -> None:
    def run() -> None:
        while True:
            try:
                run_scheduled_backup(application)
            except (OSError, sqlite3.Error):
                application.logger.exception("Scheduled ColdTrack database backup failed.")
            time.sleep(30)

    threading.Thread(target=run, name="coldtrack-backup-scheduler", daemon=True).start()


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def clean_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return " ".join(str(value).split()).upper()


def parse_iso_date(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    try:
        return date.fromisoformat(text).isoformat()
    except ValueError as exc:
        raise ValueError("Dates must use YYYY-MM-DD format.") from exc


def json_value(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def audit(
    db: sqlite3.Connection,
    actor: str,
    user_id: int | None,
    action: str,
    entity_type: str,
    entity_id: Any,
    before: Any = None,
    after: Any = None,
) -> None:
    db.execute(
        """INSERT INTO audit_log
           (user_id, actor, action, entity_type, entity_id, before_json, after_json, created_at)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
        (
            user_id,
            actor,
            action,
            entity_type,
            str(entity_id),
            json_value(before) if before is not None else None,
            json_value(after) if after is not None else None,
            now_iso(),
        ),
    )


def add_lookup(
    db: sqlite3.Connection,
    field: str,
    value: Any,
    source: str = "import",
    created_by: int | None = None,
    parent_value: str = "",
) -> None:
    cleaned = clean_value(value)
    if not cleaned or field not in LOOKUP_FIELDS:
        return
    db.execute(
        """INSERT INTO lookup_values (field, value, parent_value, source, created_by)
           VALUES (?, ?, ?, ?, ?)
           ON CONFLICT(field, value) DO UPDATE SET active = 1""",
        (field, cleaned, clean_value(parent_value), source, created_by),
    )


def store_identity(values: dict[str, str]) -> tuple[str, str, str, str]:
    return (
        clean_value(values.get("store_name")),
        " ".join(values.get("business_address", "").split()).upper(),
        clean_value(values.get("area")),
        clean_value(values.get("province")),
    )


def ensure_store(db: sqlite3.Connection, values: dict[str, str]) -> int:
    name, address, area, province = store_identity(values)
    store = db.execute(
        """SELECT id FROM stores
           WHERE store_name=? AND business_address=? AND area=? AND province=?""",
        (name, address, area, province),
    ).fetchone()
    if store:
        db.execute(
            """UPDATE stores SET manager=COALESCE(NULLIF(?, ''), manager),
               contact=COALESCE(NULLIF(?, ''), contact),
               sales_rep=COALESCE(NULLIF(?, ''), sales_rep),
               updated_at=? WHERE id=?""",
            (
                clean_value(values.get("manager")),
                " ".join(values.get("contact", "").split()),
                clean_value(values.get("sales_rep")),
                now_iso(),
                store["id"],
            ),
        )
        return store["id"]
    timestamp = now_iso()
    cursor = db.execute(
        """INSERT INTO stores
           (store_name, business_address, area, province, manager, contact, sales_rep, notes, created_at, updated_at)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (
            name, address, area, province,
            clean_value(values.get("manager")),
            " ".join(values.get("contact", "").split()),
            clean_value(values.get("sales_rep")),
            str(values.get("store_notes", values.get("notes", "")) or "").strip(),
            timestamp, timestamp,
        ),
    )
    return cursor.lastrowid


def migrate_store_data(db: sqlite3.Connection) -> None:
    store_columns = {row["name"] for row in db.execute("PRAGMA table_info(stores)")}
    if "sales_rep" not in store_columns:
        db.execute("ALTER TABLE stores ADD COLUMN sales_rep TEXT NOT NULL DEFAULT ''")
    if "notes" not in store_columns:
        db.execute("ALTER TABLE stores ADD COLUMN notes TEXT NOT NULL DEFAULT ''")
    columns = {row["name"] for row in db.execute("PRAGMA table_info(assets)")}
    if "store_id" not in columns:
        db.execute("ALTER TABLE assets ADD COLUMN store_id INTEGER REFERENCES stores(id)")
    if "moved_from_asset_id" not in columns:
        db.execute("ALTER TABLE assets ADD COLUMN moved_from_asset_id INTEGER REFERENCES assets(id)")
    if "placement_active" not in columns:
        db.execute("ALTER TABLE assets ADD COLUMN placement_active INTEGER NOT NULL DEFAULT 1")
    for asset in db.execute(
        """SELECT id, store_name, business_address, area, province, manager, contact
           FROM assets WHERE store_id IS NULL AND store_name != ''"""
    ).fetchall():
        store_id = ensure_store(db, dict(asset))
        db.execute("UPDATE assets SET store_id=? WHERE id=?", (store_id, asset["id"]))
    db.execute(
        """UPDATE stores SET sales_rep=(
               SELECT a.sales_rep FROM assets a
               WHERE a.store_id=stores.id AND a.sales_rep!=''
               ORDER BY a.updated_at DESC, a.id DESC LIMIT 1
           ) WHERE sales_rep='' AND EXISTS (
               SELECT 1 FROM assets a WHERE a.store_id=stores.id AND a.sales_rep!=''
           )"""
    )


def create_app(test_config: dict[str, Any] | None = None) -> Flask:
    app = Flask(__name__)
    INSTANCE.mkdir(parents=True, exist_ok=True)
    secret_path = INSTANCE / ".secret_key"
    if not secret_path.exists():
        secret_path.write_text(secrets.token_urlsafe(48), encoding="utf-8")
    app.config.update(
        SECRET_KEY=os.environ.get("APP_SECRET_KEY") or secret_path.read_text(encoding="utf-8").strip(),
        DATABASE=str(DEFAULT_DATABASE),
        BACKUP_DIR=str(INSTANCE / "backups"),
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
        MAX_CONTENT_LENGTH=2 * 1024 * 1024,
    )
    if test_config:
        app.config.update(test_config)

    def get_db() -> sqlite3.Connection:
        if "db" not in g:
            g.db = sqlite3.connect(app.config["DATABASE"])
            g.db.row_factory = sqlite3.Row
            g.db.execute("PRAGMA foreign_keys = ON")
            g.db.execute("PRAGMA journal_mode = WAL")
        return g.db

    app.get_db = get_db  # type: ignore[attr-defined]

    @app.teardown_appcontext
    def close_db(_error: BaseException | None = None) -> None:
        db = g.pop("db", None)
        if db is not None:
            db.close()

    with app.app_context():
        db = get_db()
        db.executescript(SCHEMA)
        db.execute(
            "INSERT OR IGNORE INTO app_settings (key, value) VALUES ('backup_time', '02:00')"
        )
        migrate_store_data(db)
        db.execute("CREATE INDEX IF NOT EXISTS assets_store_idx ON assets(store_id, placement_active)")
        seeded_reference = db.execute(
            "SELECT value FROM app_settings WHERE key='location_reference_seeded'"
        ).fetchone()
        if not seeded_reference:
            seed_location_reference(db, ROOT / "Developers.xlsx")
            db.execute(
                "INSERT INTO app_settings (key, value) VALUES ('location_reference_seeded', '1')"
            )
        db.commit()

    def user_required(handler):
        @wraps(handler)
        def wrapped(*args, **kwargs):
            if not session.get("user_id"):
                return jsonify(error="Sign in to continue."), 401
            user = get_db().execute(
                "SELECT id, username, display_name, role, active FROM users WHERE id = ?",
                (session["user_id"],),
            ).fetchone()
            if not user or not user["active"]:
                session.clear()
                return jsonify(error="This account is inactive. Contact an administrator."), 401
            g.user = user
            return handler(*args, **kwargs)

        return wrapped

    def admin_required(handler):
        @wraps(handler)
        @user_required
        def wrapped(*args, **kwargs):
            if g.user["role"] != "admin":
                return jsonify(error="Administrator access is required."), 403
            return handler(*args, **kwargs)

        return wrapped

    @app.before_request
    def protect_mutations():
        if request.method in {"POST", "PUT", "PATCH", "DELETE"} and request.path != "/api/login":
            expected = session.get("csrf_token")
            supplied = request.headers.get("X-CSRF-Token", "")
            if not expected or not secrets.compare_digest(expected, supplied):
                return jsonify(error="The request token is missing or expired. Refresh and try again."), 403

    @app.get("/")
    def index():
        return render_template("index.html")

    @app.get("/api/me")
    def me():
        if not session.get("csrf_token"):
            session["csrf_token"] = secrets.token_urlsafe(32)
        user = None
        if session.get("user_id"):
            user = get_db().execute(
                "SELECT id, username, display_name, role, active FROM users WHERE id = ? AND active = 1",
                (session["user_id"],),
            ).fetchone()
        return jsonify(user=dict(user) if user else None, csrf_token=session["csrf_token"])

    @app.post("/api/login")
    def login():
        payload = request.get_json(silent=True) or {}
        username = str(payload.get("username", "")).strip()
        password = str(payload.get("password", ""))
        user = get_db().execute(
            "SELECT * FROM users WHERE username = ? COLLATE NOCASE AND active = 1",
            (username,),
        ).fetchone()
        if not user or not check_password_hash(user["password_hash"], password):
            return jsonify(error="Username or password is incorrect."), 401
        session.clear()
        session["user_id"] = user["id"]
        session["csrf_token"] = secrets.token_urlsafe(32)
        return jsonify(
            user={
                "id": user["id"],
                "username": user["username"],
                "display_name": user["display_name"],
                "role": user["role"],
            },
            csrf_token=session["csrf_token"],
        )

    @app.post("/api/logout")
    @user_required
    def logout():
        session.clear()
        return jsonify(ok=True)

    @app.get("/api/lookups")
    @user_required
    def lookups():
        rows = get_db().execute(
            """SELECT lv.id, lv.field, lv.value, lv.parent_value FROM lookup_values lv
               WHERE lv.active = 1 AND (
                   lv.field != 'province' OR EXISTS (
                       SELECT 1 FROM lookup_values a
                       WHERE a.field = 'area' AND a.value = lv.parent_value AND a.active = 1
                   )
               )
               ORDER BY lv.field, lv.value"""
        ).fetchall()
        values: dict[str, list[dict[str, Any]]] = {field: [] for field in LOOKUP_FIELDS}
        for row in rows:
            values[row["field"]].append(
                {"id": row["id"], "value": row["value"], "parent_value": row["parent_value"]}
            )
        return jsonify(values)

    def normalize_store_values(payload: dict[str, Any], existing: sqlite3.Row | None = None) -> dict[str, str]:
        values = {
            field: (existing[field] if existing is not None else "")
            for field in (
                "store_name", "business_address", "area", "province", "manager", "contact",
                "sales_rep", "notes",
            )
        }
        for field in values:
            if field in payload:
                values[field] = str(payload[field] or "").strip()
        values["store_name"] = clean_value(values["store_name"])
        values["business_address"] = " ".join(values["business_address"].split()).upper()
        values["area"] = clean_value(values["area"])
        values["province"] = clean_value(values["province"])
        values["manager"] = clean_value(values["manager"])
        values["contact"] = " ".join(values["contact"].split())
        values["sales_rep"] = clean_value(values["sales_rep"])
        values["notes"] = str(values["notes"] or "").strip()
        if not values["store_name"]:
            raise ValueError("Store / customer name is required.")
        if not values["business_address"]:
            raise ValueError("Business address is required to identify the store.")
        db = get_db()
        if values["province"]:
            province = db.execute(
                "SELECT parent_value FROM lookup_values WHERE field='province' AND value=? AND active=1",
                (values["province"],),
            ).fetchone()
            if not province:
                raise ValueError("Choose a province from the approved list.")
            if values["area"] and values["area"] != province["parent_value"]:
                raise ValueError("The selected province does not belong to the selected area.")
            values["area"] = province["parent_value"]
        elif values["area"]:
            raise ValueError("Choose a province within the selected area.")
        return values

    @app.get("/api/stores")
    @user_required
    def list_stores():
        search = request.args.get("search", "").strip().lower()
        where = ""
        params: list[Any] = []
        if search:
            where = """WHERE lower(s.store_name) LIKE ? OR lower(s.business_address) LIKE ?
                       OR lower(s.area) LIKE ? OR lower(s.province) LIKE ?
                       OR lower(s.manager) LIKE ? OR lower(s.contact) LIKE ?
                       OR lower(s.sales_rep) LIKE ? OR lower(s.notes) LIKE ?"""
            params.extend([f"%{search}%"] * 8)
        rows = get_db().execute(
            f"""SELECT s.*,
                       COUNT(DISTINCT a.id) AS freezer_count,
                       COUNT(se.id) AS service_count,
                       MAX(se.service_date) AS latest_service
                FROM stores s
                LEFT JOIN assets a ON a.store_id=s.id AND a.placement_active=1
                LEFT JOIN service_events se ON se.asset_id=a.id
                {where}
                GROUP BY s.id
                ORDER BY s.store_name, s.business_address, s.id""",
            params,
        ).fetchall()
        return jsonify(items=[dict(row) for row in rows], total=len(rows))

    @app.get("/api/stores/<int:store_id>")
    @user_required
    def store_detail(store_id: int):
        db = get_db()
        store = db.execute("SELECT * FROM stores WHERE id=?", (store_id,)).fetchone()
        if not store:
            return jsonify(error="Store / customer not found."), 404
        freezers = db.execute(
            """SELECT a.*,
                      (SELECT COUNT(*) FROM service_events se WHERE se.asset_id=a.id) AS service_count,
                      (SELECT MAX(se.service_date) FROM service_events se WHERE se.asset_id=a.id) AS latest_service,
                      (SELECT se.technician FROM service_events se WHERE se.asset_id=a.id
                       ORDER BY se.service_date DESC, se.id DESC LIMIT 1) AS latest_technician,
                      (SELECT trim(
                          COALESCE(NULLIF(se.repairs, ''), '') ||
                          CASE WHEN se.repairs != '' AND se.parts != '' THEN ' · ' ELSE '' END ||
                          COALESCE(NULLIF(se.parts, ''), '') ||
                          CASE WHEN (se.repairs != '' OR se.parts != '') AND se.notes != '' THEN ' · ' ELSE '' END ||
                          COALESCE(NULLIF(se.notes, ''), '')
                       ) FROM service_events se WHERE se.asset_id=a.id
                       ORDER BY se.service_date DESC, se.id DESC LIMIT 1) AS latest_service_details
               FROM assets a
               WHERE a.store_id=? AND a.placement_active=1
               ORDER BY a.product_line, a.equipment_brand, a.model, a.serial_number, a.id""",
            (store_id,),
        ).fetchall()
        return jsonify(store=dict(store), freezers=[dict(row) for row in freezers])

    @app.patch("/api/stores/<int:store_id>")
    @user_required
    def update_store(store_id: int):
        payload = request.get_json(silent=True) or {}
        db = get_db()
        existing = db.execute("SELECT * FROM stores WHERE id=?", (store_id,)).fetchone()
        if not existing:
            return jsonify(error="Store / customer not found."), 404
        try:
            values = normalize_store_values(payload, existing)
        except ValueError as exc:
            return jsonify(error=str(exc)), 400
        identity = store_identity(values)
        duplicate = db.execute(
            """SELECT id FROM stores
               WHERE store_name=? AND business_address=? AND area=? AND province=? AND id!=?""",
            (*identity, store_id),
        ).fetchone()
        if duplicate:
            return jsonify(error="A store with this name and address already exists."), 409
        before = dict(existing)
        timestamp = now_iso()
        db.execute(
            """UPDATE stores SET store_name=?, business_address=?, area=?, province=?,
               manager=?, contact=?, sales_rep=?, notes=?, updated_at=? WHERE id=?""",
            (
                values["store_name"], values["business_address"], values["area"], values["province"],
                values["manager"], values["contact"], values["sales_rep"], values["notes"], timestamp, store_id,
            ),
        )
        db.execute(
            """UPDATE assets SET store_name=?, business_address=?, area=?, province=?, manager=?, contact=?, sales_rep=?,
               updated_at=? WHERE store_id=?""",
            (
                values["store_name"], values["business_address"], values["area"], values["province"],
                values["manager"], values["contact"], values["sales_rep"], timestamp, store_id,
            ),
        )
        audit(db, g.user["display_name"], g.user["id"], "update", "store", store_id, before, values)
        db.commit()
        return jsonify(ok=True)

    @app.post("/api/services")
    @user_required
    def create_service_entry():
        payload = request.get_json(silent=True) or {}
        db = get_db()
        try:
            if payload.get("new_store"):
                values = normalize_store_values(payload["new_store"])
                store_id = ensure_store(db, values)
            else:
                store_id = int(payload.get("store_id", 0))
                if not db.execute("SELECT 1 FROM stores WHERE id=?", (store_id,)).fetchone():
                    return jsonify(error="Choose an existing store or enter the new store details."), 400
            if payload.get("new_freezer"):
                store = db.execute("SELECT * FROM stores WHERE id=?", (store_id,)).fetchone()
                freezer_payload = {
                    **payload["new_freezer"],
                    "store_name": store["store_name"],
                    "business_address": store["business_address"],
                    "area": store["area"],
                    "province": store["province"],
                    "manager": store["manager"],
                    "contact": store["contact"],
                    "sales_rep": store["sales_rep"],
                    "status": "deployed",
                }
                asset_values = asset_payload(freezer_payload)
                timestamp = now_iso()
                cursor = db.execute(
                    """INSERT INTO assets
                       (store_id, product_line, equipment_type, equipment_kind, equipment_brand, model, serial_number,
                        store_name, business_address, area, province, manager, contact, delivery_date, sales_rep,
                        status, notes, created_at, updated_at)
                       VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                    (
                        store_id, asset_values["product_line"], asset_values["equipment_type"],
                        "chiller" if "CHILLER" in asset_values["equipment_type"] else "freezer",
                        asset_values["equipment_brand"], asset_values["model"], asset_values["serial_number"],
                        asset_values["store_name"], asset_values["business_address"], asset_values["area"],
                        asset_values["province"], asset_values["manager"], asset_values["contact"],
                        asset_values["delivery_date"], asset_values["sales_rep"], asset_values["status"],
                        asset_values["notes"], timestamp, timestamp,
                    ),
                )
                asset_id = cursor.lastrowid
                add_asset_lookups(db, asset_values, g.user["id"])
                db.execute(
                    """INSERT INTO location_events
                       (asset_id, business_address, area, province, effective_date, notes, created_by, created_at)
                       VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
                    (
                        asset_id, asset_values["business_address"], asset_values["area"], asset_values["province"],
                        parse_iso_date(payload.get("service_date")) or date.today().isoformat(),
                        "Initial placement created with the first recorded service.", g.user["id"], timestamp,
                    ),
                )
                audit(db, g.user["display_name"], g.user["id"], "create", "asset", asset_id, after=asset_values)
            else:
                asset_id = int(payload.get("asset_id", 0))
                freezer = db.execute(
                    "SELECT id FROM assets WHERE id=? AND store_id=? AND placement_active=1",
                    (asset_id, store_id),
                ).fetchone()
                if not freezer:
                    return jsonify(error="Choose a freezer currently listed at the selected store."), 400
            service_date = parse_iso_date(payload.get("service_date"))
            technician = clean_value(payload.get("technician"))
            if not service_date or not technician:
                return jsonify(error="Service date and technician are required."), 400
            if not db.execute(
                "SELECT 1 FROM lookup_values WHERE field='technician' AND value=? AND active=1",
                (technician,),
            ).fetchone():
                return jsonify(error="Choose a technician from the managed list."), 400
        except (TypeError, ValueError) as exc:
            return jsonify(error=str(exc)), 400
        service_values = {
            "service_date": service_date,
            "technician": technician,
            "repairs": str(payload.get("repairs", "")).strip(),
            "parts": str(payload.get("parts", "")).strip(),
            "notes": str(payload.get("notes", "")).strip(),
        }
        cursor = db.execute(
            """INSERT INTO service_events
               (asset_id, service_date, technician, repairs, parts, notes, created_by, created_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                asset_id, service_values["service_date"], service_values["technician"],
                service_values["repairs"], service_values["parts"], service_values["notes"],
                g.user["id"], now_iso(),
            ),
        )
        audit(
            db, g.user["display_name"], g.user["id"], "create", "service",
            cursor.lastrowid, after={"asset_id": asset_id, **service_values},
        )
        db.commit()
        return jsonify(id=cursor.lastrowid, asset_id=asset_id, store_id=store_id), 201

    @app.get("/api/assets")
    @user_required
    def list_assets():
        db = get_db()
        where: list[str] = []
        params: list[Any] = []
        search = request.args.get("search", "").strip()
        if search:
            like = f"%{search.lower()}%"
            where.append(
                """(
                    lower(a.product_line) LIKE ? OR lower(a.source_sheet) LIKE ? OR
                    lower(a.equipment_type) LIKE ? OR lower(a.equipment_brand) LIKE ? OR
                    lower(a.model) LIKE ? OR lower(a.serial_number) LIKE ? OR
                    lower(a.store_name) LIKE ? OR lower(a.business_address) LIKE ? OR
                    lower(a.area) LIKE ? OR lower(a.province) LIKE ? OR lower(a.manager) LIKE ? OR
                    lower(a.contact) LIKE ? OR lower(a.sales_rep) LIKE ? OR
                    lower(a.notes) LIKE ? OR
                    EXISTS (SELECT 1 FROM service_events se WHERE se.asset_id = a.id AND
                        (lower(se.technician) LIKE ? OR lower(se.repairs) LIKE ? OR lower(se.parts) LIKE ?))
                )"""
            )
            params.extend([like] * 17)
        for field, column in FILTER_FIELDS.items():
            value = request.args.get(field, "").strip()
            if value:
                if field in {"serial_number", "contact", "business_address"}:
                    where.append(f"lower({column}) LIKE ?")
                    params.append(f"%{value.lower()}%")
                else:
                    where.append(f"{column} = ?")
                    params.append(value)
        if request.args.get("current_only") == "1":
            where.append("a.placement_active=1")
        if request.args.get("available") == "1":
            where.extend(("a.placement_active=1", "a.store_id IS NULL", "a.status='unassigned'"))
        technician = request.args.get("technician", "").strip()
        if technician:
            where.append(
                "EXISTS (SELECT 1 FROM service_events st WHERE st.asset_id = a.id AND st.technician = ?)"
            )
            params.append(technician)
        requested_ids = request.args.get("ids", "")
        if requested_ids:
            try:
                asset_ids = [int(value) for value in requested_ids.split(",") if value]
            except ValueError:
                return jsonify(error="Invalid equipment selection."), 400
            if not asset_ids or len(asset_ids) > 500:
                return jsonify(error="Select between 1 and 500 equipment records."), 400
            where.append(f"a.id IN ({', '.join('?' for _ in asset_ids)})")
            params.extend(asset_ids)
        for field in ("product_line", "store_name", "equipment_type", "equipment_brand", "model"):
            value = request.args.get(f"group_{field}", "")
            if value:
                column = FILTER_FIELDS[field]
                where.append(f"{column} = ?")
                params.append(value)
        clause = f"WHERE {' AND '.join(where)}" if where else ""
        page = max(1, request.args.get("page", default=1, type=int) or 1)
        per_page = min(
            500 if request.args.get("ids") or request.args.get("available") == "1" else 100,
            max(10, request.args.get("per_page", default=50, type=int) or 50),
        )
        sort_field = request.args.get("sort", "store_name")
        sort_column = SORT_FIELDS.get(sort_field, SORT_FIELDS["store_name"])
        direction = "DESC" if request.args.get("direction", "asc").lower() == "desc" else "ASC"
        if request.args.get("grouped") == "1":
            group_columns = (
                "a.product_line",
                "a.store_name",
                "a.equipment_type",
                "a.equipment_brand",
                "a.model",
                "a.business_address",
                "a.area",
                "a.province",
            )
            group_by = ", ".join(group_columns)
            grouped_sort = (
                sort_column
                if sort_field in {"product_line", "store_name", "equipment_type", "equipment_brand", "model"}
                else "a.store_name"
            )
            total = db.execute(
                f"SELECT COUNT(*) FROM (SELECT 1 FROM assets a {clause} GROUP BY {group_by})",
                params,
            ).fetchone()[0]
            rows = db.execute(
                f"""SELECT a.product_line, a.store_name, a.equipment_type, a.equipment_brand, a.model,
                           a.business_address, a.area, a.province, COUNT(*) AS unit_count,
                           SUM((SELECT COUNT(*) FROM service_events se WHERE se.asset_id = a.id)) AS service_count,
                           GROUP_CONCAT(a.id) AS member_ids,
                           MIN(a.id) AS id
                    FROM assets a {clause}
                    GROUP BY {group_by}
                    ORDER BY {grouped_sort} {direction}, a.store_name, a.product_line, a.equipment_type, a.model
                    LIMIT ? OFFSET ?""",
                [*params, per_page, (page - 1) * per_page],
            ).fetchall()
            return jsonify(
                items=[dict(row) for row in rows],
                total=total,
                total_units=db.execute(f"SELECT COUNT(*) FROM assets a {clause}", params).fetchone()[0],
                page=page,
                per_page=per_page,
            )
        total = db.execute(f"SELECT COUNT(*) FROM assets a {clause}", params).fetchone()[0]
        rows = db.execute(
            f"""SELECT a.*,
                    (SELECT COUNT(*) FROM service_events se WHERE se.asset_id = a.id) AS service_count
                FROM assets a {clause}
                ORDER BY {sort_column} {direction}, a.store_name, a.serial_number, a.id
                LIMIT ? OFFSET ?""",
            [*params, per_page, (page - 1) * per_page],
        ).fetchall()
        return jsonify(items=[dict(row) for row in rows], total=total, page=page, per_page=per_page)

    @app.get("/api/assets/<int:asset_id>")
    @user_required
    def asset_detail(asset_id: int):
        db = get_db()
        asset = db.execute("SELECT * FROM assets WHERE id = ?", (asset_id,)).fetchone()
        if not asset:
            return jsonify(error="Equipment record not found."), 404
        services = db.execute(
            """SELECT se.*, u.display_name AS created_by_name
               FROM service_events se LEFT JOIN users u ON u.id = se.created_by
               WHERE se.asset_id = ? ORDER BY se.service_date DESC, se.id DESC""",
            (asset_id,),
        ).fetchall()
        locations = db.execute(
            """SELECT le.*, u.display_name AS created_by_name
               FROM location_events le LEFT JOIN users u ON u.id = le.created_by
               WHERE le.asset_id = ? ORDER BY COALESCE(le.effective_date, '') DESC, le.id DESC""",
            (asset_id,),
        ).fetchall()
        return jsonify(asset=dict(asset), services=[dict(row) for row in services], locations=[dict(row) for row in locations])

    @app.post("/api/assets/<int:asset_id>/move")
    @user_required
    def move_asset_to_store(asset_id: int):
        payload = request.get_json(silent=True) or {}
        db = get_db()
        asset = db.execute(
            "SELECT * FROM assets WHERE id=? AND placement_active=1", (asset_id,)
        ).fetchone()
        if not asset:
            return jsonify(error="Current freezer placement not found."), 404
        try:
            if payload.get("new_store"):
                store_values = normalize_store_values(payload["new_store"])
                store_id = ensure_store(db, store_values)
            else:
                store_id = int(payload.get("store_id", 0))
                if not db.execute("SELECT 1 FROM stores WHERE id=?", (store_id,)).fetchone():
                    return jsonify(error="Choose a destination store."), 400
            effective_date = parse_iso_date(payload.get("effective_date"))
        except (TypeError, ValueError) as exc:
            return jsonify(error=str(exc)), 400
        if not effective_date:
            return jsonify(error="Move date is required."), 400
        if store_id == asset["store_id"]:
            return jsonify(error="Choose a different destination store."), 400

        timestamp = now_iso()
        db.execute(
            "UPDATE assets SET placement_active=0, updated_at=? WHERE id=?",
            (timestamp, asset_id),
        )
        store = db.execute("SELECT * FROM stores WHERE id=?", (store_id,)).fetchone()
        cursor = db.execute(
            """INSERT INTO assets
               (store_id, moved_from_asset_id, placement_active, product_line, equipment_type, equipment_kind,
                equipment_brand, model, serial_number, store_name, business_address, area, province, manager,
                contact, delivery_date, sales_rep, status, notes, created_at, updated_at)
               VALUES (?, ?, 1, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                store_id, asset_id, asset["product_line"], asset["equipment_type"], asset["equipment_kind"],
                asset["equipment_brand"], asset["model"], asset["serial_number"], store["store_name"],
                store["business_address"], store["area"], store["province"], store["manager"], store["contact"],
                asset["delivery_date"], store["sales_rep"], "deployed", asset["notes"], timestamp, timestamp,
            ),
        )
        new_asset_id = cursor.lastrowid
        db.execute(
            """INSERT INTO location_events
               (asset_id, business_address, area, province, effective_date, notes, created_by, created_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                new_asset_id, store["business_address"], store["area"], store["province"], effective_date,
                f"Placement created by moving freezer from placement {asset_id}.",
                g.user["id"], timestamp,
            ),
        )
        audit(
            db, g.user["display_name"], g.user["id"], "move", "asset", asset_id,
            before={"store_id": asset["store_id"], "placement_active": 1},
            after={"store_id": store_id, "new_placement_id": new_asset_id, "effective_date": effective_date},
        )
        db.commit()
        return jsonify(id=new_asset_id, store_id=store_id, previous_placement_id=asset_id), 201

    def asset_payload(payload: dict[str, Any], existing: sqlite3.Row | None = None) -> dict[str, str]:
        result = {key: (existing[key] if existing is not None else "") for key in ASSET_FIELDS}
        for key in ASSET_FIELDS:
            if key in payload:
                result[key] = str(payload[key] or "").strip()
        for key in (
            "product_line",
            "equipment_type",
            "equipment_brand",
            "model",
            "serial_number",
            "store_name",
            "area",
            "province",
            "manager",
            "sales_rep",
        ):
            result[key] = clean_value(result[key])
        result["business_address"] = " ".join(result["business_address"].split())
        result["contact"] = " ".join(result["contact"].split())
        result["notes"] = result["notes"].strip()
        result["delivery_date"] = parse_iso_date(result["delivery_date"])
        if result["status"] not in {"deployed", "in_transit", "unassigned"}:
            raise ValueError("Status must be deployed, in transit, or unassigned.")
        if not result["equipment_type"]:
            raise ValueError("Equipment description/type is required.")
        db = get_db()
        for field in LOOKUP_FIELDS - {"technician", "store_name", "manager", "sales_rep"}:
            value = result.get(field, "")
            if not value or (existing is not None and value == existing[field]):
                continue
            if not db.execute(
                "SELECT 1 FROM lookup_values WHERE field = ? AND value = ? AND active = 1",
                (field, value),
            ).fetchone():
                raise ValueError(f"Choose {field.replace('_', ' ')} from the managed list.")
        if result["province"]:
            province = db.execute(
                "SELECT parent_value FROM lookup_values WHERE field = 'province' AND value = ? AND active = 1",
                (result["province"],),
            ).fetchone()
            if not province:
                raise ValueError("Choose a province from the approved list.")
            if not result["area"]:
                result["area"] = province["parent_value"]
            elif province["parent_value"] != result["area"]:
                raise ValueError("The selected province does not belong to the selected area.")
        elif result["area"]:
            raise ValueError("Choose a province within the selected area.")
        return result

    def add_asset_lookups(db: sqlite3.Connection, values: dict[str, str], user_id: int) -> None:
        for field in (
            "product_line",
            "equipment_type",
            "equipment_brand",
            "model",
            "area",
            "province",
        ):
            add_lookup(db, field, values[field], source="entry", created_by=user_id)

    def insert_asset_record(
        db: sqlite3.Connection,
        values: dict[str, str],
        store_id: int | None,
        user_id: int,
        location_effective_date: str | None = None,
        location_notes: str = "",
    ) -> int:
        timestamp = now_iso()
        cursor = db.execute(
            """INSERT INTO assets
               (store_id, product_line, equipment_type, equipment_kind, equipment_brand, model, serial_number,
                store_name, business_address, area, province, manager, contact, delivery_date, sales_rep, status,
                notes, created_at, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                store_id, values["product_line"], values["equipment_type"],
                "chiller" if "CHILLER" in values["equipment_type"] else "freezer",
                values["equipment_brand"], values["model"], values["serial_number"], values["store_name"],
                values["business_address"], values["area"], values["province"], values["manager"], values["contact"],
                values["delivery_date"], values["sales_rep"], values["status"], values["notes"], timestamp, timestamp,
            ),
        )
        asset_id = cursor.lastrowid
        add_asset_lookups(db, values, user_id)
        if location_effective_date is not None:
            db.execute(
                """INSERT INTO location_events
                   (asset_id, business_address, area, province, effective_date, notes, created_by, created_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    asset_id, values["business_address"], values["area"], values["province"],
                    location_effective_date, location_notes, user_id, timestamp,
                ),
            )
        return asset_id

    @app.post("/api/assets")
    @user_required
    def create_asset():
        payload = request.get_json(silent=True) or {}
        try:
            values = asset_payload(payload)
        except ValueError as exc:
            return jsonify(error=str(exc)), 400
        db = get_db()
        store_id = ensure_store(db, values) if values["store_name"] else None
        location_date = None
        if values["business_address"] or values["province"]:
            try:
                location_date = parse_iso_date(payload.get("location_effective_date")) or date.today().isoformat()
            except ValueError as exc:
                return jsonify(error=str(exc)), 400
        asset_id = insert_asset_record(
            db,
            values,
            store_id,
            g.user["id"],
            location_date,
            "Initial location recorded when equipment was added.",
        )
        audit(db, g.user["display_name"], g.user["id"], "create", "asset", asset_id, after=values)
        db.commit()
        return jsonify(id=asset_id), 201

    @app.post("/api/stores/<int:store_id>/freezers")
    @user_required
    def add_freezer_to_store(store_id: int):
        payload = request.get_json(silent=True) or {}
        db = get_db()
        store = db.execute("SELECT * FROM stores WHERE id=?", (store_id,)).fetchone()
        if not store:
            return jsonify(error="Store / customer not found."), 404
        timestamp = now_iso()
        asset_id = payload.get("asset_id")
        if asset_id:
            try:
                source_id = int(asset_id)
            except (TypeError, ValueError):
                return jsonify(error="Choose available equipment to add to this store."), 400
            source = db.execute(
                """SELECT * FROM assets
                   WHERE id=? AND store_id IS NULL AND status='unassigned' AND placement_active=1""",
                (source_id,),
            ).fetchone()
            if not source:
                return jsonify(error="That equipment is no longer available in unassigned stock."), 409
            db.execute(
                "UPDATE assets SET placement_active=0, updated_at=? WHERE id=?",
                (timestamp, source_id),
            )
            cursor = db.execute(
                """INSERT INTO assets
                   (store_id, moved_from_asset_id, placement_active, product_line, equipment_type, equipment_kind,
                    equipment_brand, model, serial_number, store_name, business_address, area, province, manager,
                    contact, delivery_date, sales_rep, status, notes, created_at, updated_at)
                   VALUES (?, ?, 1, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'deployed', ?, ?, ?)""",
                (
                    store_id, source_id, source["product_line"], source["equipment_type"], source["equipment_kind"],
                    source["equipment_brand"], source["model"], source["serial_number"], store["store_name"],
                    store["business_address"], store["area"], store["province"], store["manager"], store["contact"],
                    source["delivery_date"], store["sales_rep"], source["notes"], timestamp, timestamp,
                ),
            )
            new_asset_id = cursor.lastrowid
            db.execute(
                """INSERT INTO location_events
                   (asset_id, business_address, area, province, effective_date, notes, created_by, created_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    new_asset_id, store["business_address"], store["area"], store["province"],
                    date.today().isoformat(), "Equipment added from unassigned stock.", g.user["id"], timestamp,
                ),
            )
            audit(
                db, g.user["display_name"], g.user["id"], "assign", "asset", source_id,
                before={"store_id": None, "placement_active": 1},
                after={"store_id": store_id, "new_placement_id": new_asset_id},
            )
            db.commit()
            return jsonify(id=new_asset_id, store_id=store_id, previous_placement_id=source_id), 201

        freezer_payload = payload.get("new_freezer")
        if not isinstance(freezer_payload, dict):
            return jsonify(error="Choose available equipment or enter the new equipment details."), 400
        try:
            values = asset_payload(
                {
                    **freezer_payload,
                    "store_name": store["store_name"],
                    "business_address": store["business_address"],
                    "area": store["area"],
                    "province": store["province"],
                    "manager": store["manager"],
                    "contact": store["contact"],
                    "sales_rep": store["sales_rep"],
                    "status": "deployed",
                }
            )
        except ValueError as exc:
            return jsonify(error=str(exc)), 400
        new_asset_id = insert_asset_record(
            db,
            values,
            store_id,
            g.user["id"],
            date.today().isoformat(),
            "Equipment added to this store without a service visit.",
        )
        audit(db, g.user["display_name"], g.user["id"], "create", "asset", new_asset_id, after=values)
        db.commit()
        return jsonify(id=new_asset_id, store_id=store_id), 201

    @app.patch("/api/assets/<int:asset_id>")
    @user_required
    def update_asset(asset_id: int):
        payload = request.get_json(silent=True) or {}
        db = get_db()
        existing = db.execute("SELECT * FROM assets WHERE id = ?", (asset_id,)).fetchone()
        if not existing:
            return jsonify(error="Equipment record not found."), 404
        try:
            values = asset_payload(payload, existing)
        except ValueError as exc:
            return jsonify(error=str(exc)), 400
        before = {field: existing[field] for field in ASSET_FIELDS}
        changed_location = any(
            values[field] != existing[field]
            for field in ("business_address", "area", "province")
        )
        timestamp = now_iso()
        store_id = ensure_store(db, values) if values["store_name"] else None
        db.execute(
            """UPDATE assets SET store_id=?, product_line=?, equipment_type=?, equipment_kind=?, equipment_brand=?,
               model=?, serial_number=?, store_name=?, business_address=?, area=?, province=?, manager=?, contact=?,
               delivery_date=?, sales_rep=?, status=?, notes=?, updated_at=? WHERE id=?""",
            (
                store_id, values["product_line"], values["equipment_type"],
                "chiller" if "CHILLER" in values["equipment_type"] else "freezer",
                values["equipment_brand"], values["model"], values["serial_number"],
                values["store_name"], values["business_address"], values["area"], values["province"], values["manager"],
                values["contact"], values["delivery_date"], values["sales_rep"], values["status"],
                values["notes"], timestamp, asset_id,
            ),
        )
        add_asset_lookups(db, values, g.user["id"])
        if changed_location:
            try:
                effective_date = parse_iso_date(payload.get("location_effective_date")) or date.today().isoformat()
            except ValueError as exc:
                return jsonify(error=str(exc)), 400
            db.execute(
                """INSERT INTO location_events
                   (asset_id, business_address, area, province, effective_date, notes, created_by, created_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    asset_id, values["business_address"], values["area"], values["province"], effective_date,
                    str(payload.get("location_notes", "")).strip(), g.user["id"], timestamp,
                ),
            )
        audit(db, g.user["display_name"], g.user["id"], "update", "asset", asset_id, before, values)
        db.commit()
        return jsonify(ok=True)

    @app.post("/api/assets/<int:asset_id>/services")
    @user_required
    def create_service(asset_id: int):
        payload = request.get_json(silent=True) or {}
        db = get_db()
        if not db.execute(
            "SELECT 1 FROM assets WHERE id = ? AND placement_active=1", (asset_id,)
        ).fetchone():
            return jsonify(error="Equipment record not found."), 404
        try:
            service_date = parse_iso_date(payload.get("service_date"))
        except ValueError as exc:
            return jsonify(error=str(exc)), 400
        technician = clean_value(payload.get("technician"))
        if not service_date or not technician:
            return jsonify(error="Service date and technician are required."), 400
        if not db.execute(
            "SELECT 1 FROM lookup_values WHERE field = 'technician' AND value = ? AND active = 1",
            (technician,),
        ).fetchone():
            return jsonify(error="Choose a technician from the managed list."), 400
        values = {
            "service_date": service_date,
            "technician": technician,
            "repairs": str(payload.get("repairs", "")).strip(),
            "parts": str(payload.get("parts", "")).strip(),
            "notes": str(payload.get("notes", "")).strip(),
        }
        cursor = db.execute(
            """INSERT INTO service_events
               (asset_id, service_date, technician, repairs, parts, notes, created_by, created_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                asset_id, values["service_date"], values["technician"], values["repairs"],
                values["parts"], values["notes"], g.user["id"], now_iso(),
            ),
        )
        audit(
            db, g.user["display_name"], g.user["id"], "create", "service",
            cursor.lastrowid, after={"asset_id": asset_id, **values},
        )
        db.commit()
        return jsonify(id=cursor.lastrowid), 201

    @app.patch("/api/services/<int:service_id>")
    @user_required
    def update_service(service_id: int):
        payload = request.get_json(silent=True) or {}
        db = get_db()
        existing = db.execute("SELECT * FROM service_events WHERE id = ?", (service_id,)).fetchone()
        if not existing:
            return jsonify(error="Service record not found."), 404
        try:
            service_date = parse_iso_date(payload.get("service_date", existing["service_date"]))
        except ValueError as exc:
            return jsonify(error=str(exc)), 400
        technician = clean_value(payload.get("technician", existing["technician"]))
        if not service_date or not technician:
            return jsonify(error="Service date and technician are required."), 400
        if technician != existing["technician"] and not db.execute(
            "SELECT 1 FROM lookup_values WHERE field = 'technician' AND value = ? AND active = 1",
            (technician,),
        ).fetchone():
            return jsonify(error="Choose a technician from the managed list."), 400
        before = {
            "service_date": existing["service_date"],
            "technician": existing["technician"],
            "repairs": existing["repairs"],
            "parts": existing["parts"],
            "notes": existing["notes"],
        }
        after = {
            "service_date": service_date,
            "technician": technician,
            "repairs": str(payload.get("repairs", existing["repairs"])).strip(),
            "parts": str(payload.get("parts", existing["parts"])).strip(),
            "notes": str(payload.get("notes", existing["notes"])).strip(),
        }
        db.execute(
            """UPDATE service_events SET service_date=?, technician=?, repairs=?, parts=?, notes=?
               WHERE id=?""",
            (
                after["service_date"], after["technician"], after["repairs"],
                after["parts"], after["notes"], service_id,
            ),
        )
        if technician != existing["technician"]:
            add_lookup(db, "technician", technician, source="entry", created_by=g.user["id"])
        audit(
            db, g.user["display_name"], g.user["id"], "update", "service",
            service_id, before, after,
        )
        db.commit()
        return jsonify(ok=True)

    @app.get("/api/reports/services")
    @user_required
    def service_report():
        where: list[str] = []
        params: list[Any] = []
        for key, column in (
            ("technician", "se.technician"),
            ("product_line", "a.product_line"),
            ("equipment_brand", "a.equipment_brand"),
            ("model", "a.model"),
        ):
            value = request.args.get(key, "").strip()
            if value:
                where.append(f"{column} = ?")
                params.append(value)
        serial_number = request.args.get("serial_number", "").strip()
        if serial_number:
            where.append("lower(a.serial_number) LIKE ?")
            params.append(f"%{serial_number.lower()}%")
        store_id = request.args.get("store_id", "").strip()
        if store_id:
            try:
                store_id_value = int(store_id)
            except ValueError:
                return jsonify(error="Choose a valid store."), 400
            where.append("a.store_id = ?")
            params.append(store_id_value)
        clause = f"WHERE {' AND '.join(where)}" if where else ""
        rows = get_db().execute(
            f"""SELECT a.id AS asset_id, se.technician, a.product_line, a.equipment_brand, a.model,
                       a.serial_number, a.store_name, COUNT(*) AS service_count,
                       MAX(se.service_date) AS latest_service
                FROM service_events se JOIN assets a ON a.id = se.asset_id
                {clause}
                GROUP BY se.technician, a.id
                ORDER BY service_count DESC, se.technician, a.serial_number""",
            params,
        ).fetchall()
        return jsonify(items=[dict(row) for row in rows], total_services=sum(row["service_count"] for row in rows))

    @app.get("/api/reports/services/<int:asset_id>")
    @user_required
    def service_report_detail(asset_id: int):
        db = get_db()
        asset = db.execute("SELECT * FROM assets WHERE id = ?", (asset_id,)).fetchone()
        if not asset:
            return jsonify(error="Equipment record not found."), 404
        if asset["serial_number"]:
            services = db.execute(
                """SELECT se.*, a.serial_number, a.store_name, a.business_address,
                          u.display_name AS created_by_name
                   FROM service_events se
                   JOIN assets a ON a.id = se.asset_id
                   LEFT JOIN users u ON u.id = se.created_by
                   WHERE lower(trim(a.serial_number)) = lower(trim(?))
                   ORDER BY se.service_date DESC, se.id DESC""",
                (asset["serial_number"],),
            ).fetchall()
        else:
            services = db.execute(
                """SELECT se.*, a.serial_number, a.store_name, a.business_address,
                          u.display_name AS created_by_name
                   FROM service_events se
                   JOIN assets a ON a.id = se.asset_id
                   LEFT JOIN users u ON u.id = se.created_by
                   WHERE se.asset_id = ?
                   ORDER BY se.service_date DESC, se.id DESC""",
                (asset_id,),
            ).fetchall()
        return jsonify(asset=dict(asset), services=[dict(row) for row in services])

    @app.get("/api/admin/lookups")
    @admin_required
    def admin_lookups():
        rows = get_db().execute(
            "SELECT id, field, value, parent_value, source, active FROM lookup_values ORDER BY field, value"
        ).fetchall()
        return jsonify(items=[dict(row) for row in rows])

    @app.post("/api/admin/lookups")
    @admin_required
    def create_lookup():
        payload = request.get_json(silent=True) or {}
        field = str(payload.get("field", "")).strip()
        value = clean_value(payload.get("value"))
        parent_value = clean_value(payload.get("parent_value"))
        if field not in LOOKUP_FIELDS or not value:
            return jsonify(error="Choose a valid list and enter a value."), 400
        db = get_db()
        if field == "province":
            if not parent_value or not db.execute(
                "SELECT 1 FROM lookup_values WHERE field = 'area' AND value = ? AND active = 1",
                (parent_value,),
            ).fetchone():
                return jsonify(error="Select an active area for this province."), 400
        elif parent_value:
            return jsonify(error="Only provinces can be assigned to an area."), 400
        try:
            cursor = db.execute(
                """INSERT INTO lookup_values (field, value, parent_value, source, created_by)
                   VALUES (?, ?, ?, 'admin', ?)""",
                (field, value, parent_value, g.user["id"]),
            )
        except sqlite3.IntegrityError:
            return jsonify(error="That value already exists in this list."), 409
        audit(
            db, g.user["display_name"], g.user["id"], "create", "lookup",
            cursor.lastrowid, after={"field": field, "value": value, "parent_value": parent_value},
        )
        db.commit()
        return jsonify(id=cursor.lastrowid), 201

    @app.patch("/api/admin/lookups/<int:lookup_id>")
    @admin_required
    def update_lookup(lookup_id: int):
        payload = request.get_json(silent=True) or {}
        db = get_db()
        row = db.execute("SELECT * FROM lookup_values WHERE id = ?", (lookup_id,)).fetchone()
        if not row:
            return jsonify(error="List value not found."), 404
        before = dict(row)
        parent_value = clean_value(payload.get("parent_value", row["parent_value"]))
        if "parent_value" in payload and row["field"] != "province":
            return jsonify(error="Only child values can be assigned to a parent."), 400
        if row["field"] == "province" and parent_value != row["parent_value"] and (
            not parent_value or not db.execute(
                "SELECT 1 FROM lookup_values WHERE field = 'area' AND value = ? AND active = 1",
                (parent_value,),
            ).fetchone()
        ):
            return jsonify(error="Select an active area for this province."), 400
        active = row["active"]
        if "active" in payload:
            if not isinstance(payload["active"], bool):
                return jsonify(error="Active must be true or false."), 400
            active = int(payload["active"])
        value = clean_value(payload.get("value", row["value"]))
        if not value:
            return jsonify(error="List value cannot be empty."), 400
        try:
            db.execute(
                "UPDATE lookup_values SET value=?, parent_value=?, active=? WHERE id=?",
                (value, parent_value, active, lookup_id),
            )
        except sqlite3.IntegrityError:
            return jsonify(error="That value already exists in this list. Merge it with the existing value instead."), 409
        if value != row["value"]:
            field_columns = {
                "product_line": (("assets", "product_line"),),
                "equipment_type": (("assets", "equipment_type"),),
                "equipment_brand": (("assets", "equipment_brand"),),
                "model": (("assets", "model"),),
                "store_name": (("stores", "store_name"), ("assets", "store_name")),
                "area": (("stores", "area"), ("assets", "area"), ("location_events", "area")),
                "province": (("stores", "province"), ("assets", "province"), ("location_events", "province")),
                "manager": (("stores", "manager"), ("assets", "manager")),
                "sales_rep": (("assets", "sales_rep"),),
                "technician": (("service_events", "technician"),),
            }
            for table, column in field_columns.get(row["field"], ()):
                db.execute(f"UPDATE {table} SET {column}=? WHERE {column}=?", (value, row["value"]))
            if row["field"] == "area":
                db.execute(
                    "UPDATE lookup_values SET parent_value=? WHERE field='province' AND parent_value=?",
                    (value, row["value"]),
                )
        if row["field"] == "province" and parent_value != row["parent_value"]:
            db.execute(
                """UPDATE stores SET area=? WHERE province=? AND area=?""",
                (parent_value, value, row["parent_value"]),
            )
            db.execute(
                """UPDATE assets SET area=? WHERE province=? AND area=?""",
                (parent_value, value, row["parent_value"]),
            )
            db.execute(
                """UPDATE location_events SET area=? WHERE province=? AND area=?""",
                (parent_value, value, row["parent_value"]),
            )
        audit(
            db, g.user["display_name"], g.user["id"], "update", "lookup",
            lookup_id, before, {"field": row["field"], "value": value, "parent_value": parent_value, "active": active},
        )
        db.commit()
        return jsonify(ok=True)

    @app.delete("/api/admin/lookups/<int:lookup_id>")
    @admin_required
    def deactivate_lookup(lookup_id: int):
        db = get_db()
        row = db.execute("SELECT * FROM lookup_values WHERE id = ?", (lookup_id,)).fetchone()
        if not row:
            return jsonify(error="List value not found."), 404
        before = dict(row)
        db.execute("UPDATE lookup_values SET active = 0 WHERE id = ?", (lookup_id,))
        audit(
            db, g.user["display_name"], g.user["id"], "deactivate", "lookup",
            lookup_id, before, {"field": row["field"], "value": row["value"], "active": 0},
        )
        db.commit()
        return jsonify(ok=True)

    @app.get("/api/admin/backups")
    @admin_required
    def list_backups():
        db = get_db()
        settings = {
            row["key"]: row["value"]
            for row in db.execute(
                "SELECT key, value FROM app_settings WHERE key IN ('backup_time', 'backup_last_run_date')"
            ).fetchall()
        }
        backup_directory = Path(app.config["BACKUP_DIR"])
        backups = []
        if backup_directory.exists():
            backups = [
                {
                    "filename": path.name,
                    "created_at": datetime.fromtimestamp(path.stat().st_mtime).astimezone().isoformat(timespec="minutes"),
                    "size": path.stat().st_size,
                }
                for path in sorted(
                    backup_directory.glob("coldtrack-*.sqlite3"),
                    key=lambda item: item.stat().st_mtime,
                    reverse=True,
                )
            ]
        return jsonify(
            backup_time=settings.get("backup_time", "02:00"),
            last_scheduled_backup=settings.get("backup_last_run_date", ""),
            backups=backups,
        )

    @app.patch("/api/admin/backups/schedule")
    @admin_required
    def update_backup_schedule():
        payload = request.get_json(silent=True) or {}
        backup_time = str(payload.get("time", ""))
        if not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", backup_time):
            return jsonify(error="Choose a valid backup time."), 400
        db = get_db()
        db.execute(
            """INSERT INTO app_settings (key, value) VALUES ('backup_time', ?)
               ON CONFLICT(key) DO UPDATE SET value=excluded.value""",
            (backup_time,),
        )
        audit(
            db, g.user["display_name"], g.user["id"], "update", "backup_schedule",
            "daily", after={"time": backup_time},
        )
        db.commit()
        return jsonify(ok=True, time=backup_time)

    @app.post("/api/admin/backups")
    @admin_required
    def create_backup():
        path = create_database_backup(app.config["DATABASE"], app.config["BACKUP_DIR"])
        db = get_db()
        audit(
            db, g.user["display_name"], g.user["id"], "create", "database_backup",
            path.name, after={"filename": path.name, "size": path.stat().st_size},
        )
        db.commit()
        return jsonify(filename=path.name), 201

    @app.get("/api/admin/backups/<filename>")
    @admin_required
    def download_backup(filename: str):
        if not re.fullmatch(r"coldtrack-\d{8}-\d{6}-[0-9a-f]{6}\.sqlite3", filename):
            return jsonify(error="Backup not found."), 404
        path = Path(app.config["BACKUP_DIR"]) / filename
        if not path.is_file():
            return jsonify(error="Backup not found."), 404
        return send_file(
            path,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.sqlite3",
        )

    @app.post("/api/admin/reset")
    @admin_required
    def reset_application_data():
        payload = request.get_json(silent=True) or {}
        if payload.get("confirmation") != "RESET":
            return jsonify(error='Type "RESET" to confirm clearing all application data.'), 400
        backup = create_database_backup(app.config["DATABASE"], app.config["BACKUP_DIR"])
        db = get_db()
        db.execute("BEGIN IMMEDIATE")
        for table in (
            "lookup_requests",
            "import_batches",
            "service_events",
            "location_events",
            "assets",
            "stores",
            "audit_log",
            "lookup_values",
        ):
            db.execute(f"DELETE FROM {table}")
        db.execute("DELETE FROM users WHERE id != ?", (g.user["id"],))
        db.commit()
        return jsonify(ok=True, backup=backup.name)

    @app.get("/api/admin/variants")
    @admin_required
    def lookup_variants():
        db = get_db()
        result: list[dict[str, Any]] = []
        for field in LOOKUP_FIELDS - {"area", "province", "store_name", "manager", "sales_rep"}:
            rows = db.execute(
                "SELECT id, value FROM lookup_values WHERE field = ? AND active = 1 ORDER BY value",
                (field,),
            ).fetchall()
            by_value = {row["value"]: row["id"] for row in rows}
            names = list(by_value)
            for left in names:
                for right in difflib.get_close_matches(left, names, n=2, cutoff=0.86):
                    if left >= right:
                        continue
                    score = difflib.SequenceMatcher(None, left, right).ratio()
                    result.append(
                        {
                            "field": field,
                            "left_id": by_value[left],
                            "left": left,
                            "right_id": by_value[right],
                            "right": right,
                            "similarity": round(score, 2),
                        }
                    )
        result.sort(key=lambda item: (-item["similarity"], item["field"], item["left"]))
        return jsonify(items=result[:200])

    @app.post("/api/admin/variants/merge")
    @admin_required
    def merge_lookup_values():
        payload = request.get_json(silent=True) or {}
        source_id = payload.get("source_id")
        target_id = payload.get("target_id")
        if not isinstance(source_id, int) or not isinstance(target_id, int) or source_id == target_id:
            return jsonify(error="Choose two different list values."), 400
        db = get_db()
        source = db.execute("SELECT * FROM lookup_values WHERE id = ? AND active = 1", (source_id,)).fetchone()
        target = db.execute("SELECT * FROM lookup_values WHERE id = ? AND active = 1", (target_id,)).fetchone()
        if not source or not target or source["field"] != target["field"]:
            return jsonify(error="The selected values are not in the same active list."), 400
        if source["field"] in {"area", "province"}:
            return jsonify(error="Area and Province values must be edited individually to preserve their hierarchy."), 400
        field_to_tables = {
            "product_line": (("assets", "product_line"),),
            "equipment_type": (("assets", "equipment_type"),),
            "equipment_brand": (("assets", "equipment_brand"),),
            "model": (("assets", "model"),),
            "store_name": (("stores", "store_name"), ("assets", "store_name")),
            "manager": (("stores", "manager"), ("assets", "manager")),
            "sales_rep": (("assets", "sales_rep"),),
            "technician": (("service_events", "technician"),),
        }
        for table, column in field_to_tables[source["field"]]:
            db.execute(
                f"UPDATE {table} SET {column} = ? WHERE {column} = ?",
                (target["value"], source["value"]),
            )
        before = {"field": source["field"], "source": source["value"], "target": target["value"]}
        db.execute("UPDATE lookup_values SET active = 0 WHERE id = ?", (source_id,))
        audit(
            db, g.user["display_name"], g.user["id"], "merge", "lookup",
            source_id, before, {"canonical_value": target["value"]},
        )
        db.commit()
        return jsonify(ok=True)

    @app.get("/api/admin/users")
    @admin_required
    def admin_users():
        rows = get_db().execute(
            "SELECT id, username, display_name, role, active, created_at FROM users ORDER BY username"
        ).fetchall()
        return jsonify(items=[dict(row) for row in rows])

    @app.post("/api/imports/preview")
    @user_required
    def preview_import():
        uploaded = request.files.get("workbook")
        if not uploaded or not uploaded.filename:
            return jsonify(error="Choose an Excel workbook to import."), 400
        if Path(uploaded.filename).suffix.lower() != ".xlsx":
            return jsonify(error="Choose an .xlsx workbook."), 400
        workbook_bytes = uploaded.read()
        if not workbook_bytes:
            return jsonify(error="The selected workbook is empty."), 400
        selected_sheets = [clean_value(value) for value in request.form.getlist("sheets")]
        try:
            available_sheets, records = workbook_records(workbook_bytes, selected_sheets, get_db())
        except (ValueError, OSError, zipfile.BadZipFile) as exc:
            return jsonify(error=str(exc)), 400
        if not selected_sheets:
            return jsonify(sheets=available_sheets)
        if not records:
            return jsonify(error="No equipment rows were found in the selected sheets."), 400
        db = get_db()
        cursor = db.execute(
            """INSERT INTO import_batches
               (file_name, sheets_json, records_json, created_by, status, created_at)
               VALUES (?, ?, ?, ?, 'review', ?)""",
            (
                Path(uploaded.filename).name,
                json_value(selected_sheets),
                json_value(records),
                g.user["id"],
                now_iso(),
            ),
        )
        batch_id = cursor.lastrowid
        db.commit()
        return jsonify(
            batch_id=batch_id,
            file_name=Path(uploaded.filename).name,
            sheets=selected_sheets,
            record_count=len(records),
            review=import_review(db, records),
            rows=import_preview_rows(records),
        ), 201

    @app.get("/api/imports/<int:batch_id>")
    @user_required
    def get_import_preview(batch_id: int):
        batch = get_db().execute("SELECT * FROM import_batches WHERE id = ?", (batch_id,)).fetchone()
        if not batch:
            return jsonify(error="Import batch not found."), 404
        if batch["created_by"] != g.user["id"] and g.user["role"] != "admin":
            return jsonify(error="You can only view your own import batches."), 403
        records = json.loads(batch["records_json"])
        return jsonify(
            batch_id=batch["id"],
            file_name=batch["file_name"],
            sheets=json.loads(batch["sheets_json"]),
            status=batch["status"],
            record_count=len(records),
            review=import_review(get_db(), records) if batch["status"] == "review" else None,
            rows=import_preview_rows(records) if batch["status"] == "review" else [],
        )

    @app.post("/api/imports/<int:batch_id>/commit")
    @user_required
    def commit_import_review(batch_id: int):
        payload = request.get_json(silent=True) or {}
        db = get_db()
        batch = db.execute("SELECT * FROM import_batches WHERE id = ?", (batch_id,)).fetchone()
        if not batch:
            return jsonify(error="Import batch not found."), 404
        if batch["created_by"] != g.user["id"] and g.user["role"] != "admin":
            return jsonify(error="You can only submit your own import batches."), 403
        if batch["status"] != "review":
            return jsonify(error="This import is no longer awaiting review."), 409
        records = json.loads(batch["records_json"])
        review = import_review(db, records)
        choices = payload.get("choices", {})
        locations = payload.get("locations", {})
        date_resolutions = payload.get("date_resolutions", {})
        if not isinstance(choices, dict) or any(
            not isinstance(field_choices, dict) for field_choices in choices.values()
        ):
            return jsonify(error="Import choices must be submitted as field/value mappings."), 400
        if not isinstance(date_resolutions, dict):
            return jsonify(error="Date resolutions must be submitted as row/value mappings."), 400
        parsed_dates: dict[str, str] = {}
        for warning in review["date_warnings"]:
            date_key = f"{warning['sheet']}:{warning['row']}"
            resolution = date_resolutions.get(date_key)
            if not isinstance(resolution, dict):
                return jsonify(error=f"Normalize or explicitly leave blank the delivery date in {warning['sheet']} row {warning['row']}."), 400
            if resolution.get("leave_blank") is True:
                if resolution.get("normalized") not in (None, ""):
                    return jsonify(error=f"Choose a date or leave it blank for {warning['sheet']} row {warning['row']}, not both."), 400
                parsed_dates[date_key] = ""
                continue
            normalized = resolution.get("normalized")
            if not isinstance(normalized, str):
                return jsonify(error=f"Enter a valid date or explicitly leave blank the delivery date in {warning['sheet']} row {warning['row']}."), 400
            try:
                parsed_dates[date_key] = parse_iso_date(normalized)
            except ValueError:
                return jsonify(error=f"Use YYYY-MM-DD for the delivery date in {warning['sheet']} row {warning['row']}."), 400
            if not parsed_dates[date_key]:
                return jsonify(error=f"Enter a valid date or explicitly leave blank the delivery date in {warning['sheet']} row {warning['row']}."), 400
        for question in review["unknown_values"]:
            field = question["field"]
            raw_value = question["value"]
            selected = choices.get(field, {}).get(raw_value)
            if not isinstance(selected, str) or not selected:
                return jsonify(error=f"Choose how to handle {field.replace('_', ' ').lower()} “{raw_value}”."), 400
            if selected == "__blank__" and field in {"area", "province"}:
                continue
            if selected == "__new__" and field in {"area", "province"}:
                return jsonify(error=f"Add new {field} values in Administration before importing."), 400
            if selected != "__new__":
                approved = db.execute(
                    "SELECT value FROM lookup_values WHERE field = ? AND value = ? AND active = 1",
                    (field, selected),
                ).fetchone()
                if not approved:
                    return jsonify(error=f"“{selected}” is not an approved {field.replace('_', ' ')} value."), 400
        if review["location_warnings"] and payload.get("acknowledge_location_warnings") is not True:
                return jsonify(
                    error="Review and acknowledge the rows with missing Area or Province values before importing."
                ), 400
        province_rows = db.execute(
            "SELECT value, parent_value FROM lookup_values WHERE field = 'province' AND active = 1"
        ).fetchall()
        province_areas = {row["value"]: row["parent_value"] for row in province_rows}
        for raw_location in review["location_questions"]:
            resolution = locations.get(raw_location)
            if not isinstance(resolution, dict):
                return jsonify(error=f"Clarify the meaning of source location “{raw_location}”."), 400
            if resolution.get("kind") == "address":
                address = " ".join(str(resolution.get("address", "")).split())
                if not address:
                    return jsonify(error=f"Enter a business address for “{raw_location}” or choose a reference province."), 400
            elif resolution.get("kind") == "province":
                area = clean_value(resolution.get("area"))
                province = clean_value(resolution.get("province"))
                if not province or province_areas.get(province) != area:
                    return jsonify(error=f"Choose a valid area and province for “{raw_location}”."), 400
            else:
                return jsonify(error=f"Choose whether “{raw_location}” is a business address or a reference province."), 400

        pending: list[tuple[str, str]] = []
        for record in records:
            for field in IMPORT_LOOKUP_FIELDS:
                raw_value = clean_value(record.get(field))
                if not raw_value:
                    record[field] = ""
                    continue
                selected = choices.get(field, {}).get(raw_value, raw_value)
                if selected == "__blank__" and field in {"area", "province"}:
                    record[field] = ""
                elif selected == "__new__":
                    record[field] = raw_value
                    pending.append((field, raw_value))
                else:
                    record[field] = selected
            source_location = record.get("source_location", "")
            date_key = f"{record['source_sheet']}:{record['source_row']}"
            if date_key in parsed_dates:
                record["delivery_date"] = parsed_dates[date_key]
            if record.get("location_needs_review"):
                resolution = locations[source_location]
                if resolution["kind"] == "address":
                    record["business_address"] = " ".join(resolution["address"].split())
                    record["area"] = ""
                    record["province"] = ""
                else:
                    record["area"] = clean_value(resolution["area"])
                    record["province"] = clean_value(resolution["province"])
            record.pop("source_location", None)
            record.pop("location_needs_review", None)
            record.pop("unparsed_delivery_date", None)
            if record.get("province"):
                parent_area = province_areas.get(record["province"])
                if not parent_area or record.get("area") != parent_area:
                    return jsonify(
                        error=(
                            f"Area and Province do not match for {record['source_sheet']} row "
                            f"{record['source_row']}. Select a matching pair."
                        )
                    ), 400
            record.pop("location_review_required", None)

        db.execute("UPDATE import_batches SET records_json=? WHERE id=?", (json_value(records), batch_id))
        for field, value in sorted(set(pending)):
            db.execute(
                """INSERT OR IGNORE INTO lookup_requests
                   (batch_id, field, value, status, requested_by, created_at)
                   VALUES (?, ?, ?, 'pending', ?, ?)""",
                (batch_id, field, value, g.user["id"], now_iso()),
            )
        pending_count = db.execute(
            "SELECT COUNT(*) FROM lookup_requests WHERE batch_id=? AND status='pending'",
            (batch_id,),
        ).fetchone()[0]
        if pending_count:
            db.execute("UPDATE import_batches SET status='awaiting_approval' WHERE id=?", (batch_id,))
            audit(
                db, g.user["display_name"], g.user["id"], "submit_import",
                "import_batch", batch_id,
                after={"file_name": batch["file_name"], "records": len(records), "pending_values": pending_count},
            )
            db.commit()
            return jsonify(status="awaiting_approval", pending_values=pending_count, records=len(records))
        apply_import_batch(db, batch, g.user)
        db.commit()
        return jsonify(status="complete", imported=len(records))

    @app.post("/api/imports/<int:batch_id>/commit-rows")
    @user_required
    def commit_import_rows(batch_id: int):
        payload = request.get_json(silent=True) or {}
        db = get_db()
        batch = db.execute("SELECT * FROM import_batches WHERE id = ?", (batch_id,)).fetchone()
        if not batch:
            return jsonify(error="Import batch not found."), 404
        if batch["created_by"] != g.user["id"] and g.user["role"] != "admin":
            return jsonify(error="You can only submit your own import batches."), 403
        if batch["status"] != "review":
            return jsonify(error="This import is no longer awaiting review."), 409

        records = json.loads(batch["records_json"])
        try:
            pending = apply_import_row_edits(
                db,
                records,
                payload.get("rows"),
                payload.get("acknowledge_location_warnings") is True,
            )
        except ValueError as exc:
            return jsonify(error=str(exc)), 400

        db.execute("UPDATE import_batches SET records_json=? WHERE id=?", (json_value(records), batch_id))
        for field, value in sorted(set(pending)):
            db.execute(
                """INSERT OR IGNORE INTO lookup_requests
                   (batch_id, field, value, status, requested_by, created_at)
                   VALUES (?, ?, ?, 'pending', ?, ?)""",
                (batch_id, field, value, g.user["id"], now_iso()),
            )
        pending_count = db.execute(
            "SELECT COUNT(*) FROM lookup_requests WHERE batch_id=? AND status='pending'",
            (batch_id,),
        ).fetchone()[0]
        if pending_count:
            db.execute("UPDATE import_batches SET status='awaiting_approval' WHERE id=?", (batch_id,))
            audit(
                db, g.user["display_name"], g.user["id"], "submit_import",
                "import_batch", batch_id,
                after={"file_name": batch["file_name"], "records": len(records), "pending_values": pending_count},
            )
            db.commit()
            return jsonify(status="awaiting_approval", pending_values=pending_count, records=len(records))
        apply_import_batch(db, batch, g.user)
        db.commit()
        return jsonify(status="complete", imported=len(records))

    @app.get("/api/admin/import-requests")
    @admin_required
    def admin_import_requests():
        rows = get_db().execute(
            """SELECT r.id, r.batch_id, r.field, r.value, r.created_at,
                      u.display_name AS requested_by, b.file_name
               FROM lookup_requests r
               JOIN import_batches b ON b.id = r.batch_id
               LEFT JOIN users u ON u.id = r.requested_by
               WHERE r.status = 'pending'
               ORDER BY r.created_at, r.field, r.value"""
        ).fetchall()
        return jsonify(items=[dict(row) for row in rows])

    @app.post("/api/admin/import-requests/<int:request_id>/approve")
    @admin_required
    def approve_import_lookup_request(request_id: int):
        db = get_db()
        lookup_request = db.execute(
            "SELECT * FROM lookup_requests WHERE id=? AND status='pending'", (request_id,)
        ).fetchone()
        if not lookup_request:
            return jsonify(error="Pending import value not found."), 404
        if lookup_request["field"] in {"area", "province"}:
            return jsonify(error="Area and province values must come from the reference workbook."), 400
        add_lookup(
            db,
            lookup_request["field"],
            lookup_request["value"],
            source="approved-import",
            created_by=g.user["id"],
        )
        timestamp = now_iso()
        db.execute(
            """UPDATE lookup_requests SET status='approved', reviewed_by=?, reviewed_at=?
               WHERE id=?""",
            (g.user["id"], timestamp, request_id),
        )
        audit(
            db, g.user["display_name"], g.user["id"], "approve", "lookup_request",
            request_id,
            after={"field": lookup_request["field"], "value": lookup_request["value"]},
        )
        batch = db.execute(
            "SELECT * FROM import_batches WHERE id=?", (lookup_request["batch_id"],)
        ).fetchone()
        pending_count = db.execute(
            "SELECT COUNT(*) FROM lookup_requests WHERE batch_id=? AND status='pending'",
            (batch["id"],),
        ).fetchone()[0]
        imported = 0
        if not pending_count:
            imported = apply_import_batch(db, batch, g.user)
        db.commit()
        return jsonify(ok=True, batch_status="complete" if not pending_count else "awaiting_approval", imported=imported)

    @app.post("/api/admin/import-requests/<int:request_id>/reject")
    @admin_required
    def reject_import_lookup_request(request_id: int):
        db = get_db()
        lookup_request = db.execute(
            "SELECT * FROM lookup_requests WHERE id=? AND status='pending'", (request_id,)
        ).fetchone()
        if not lookup_request:
            return jsonify(error="Pending import value not found."), 404
        timestamp = now_iso()
        db.execute(
            """UPDATE lookup_requests SET status='rejected', reviewed_by=?, reviewed_at=?
               WHERE id=?""",
            (g.user["id"], timestamp, request_id),
        )
        db.execute(
            "UPDATE import_batches SET status='rejected', completed_at=? WHERE id=?",
            (timestamp, lookup_request["batch_id"]),
        )
        db.execute(
            """UPDATE lookup_requests SET status='rejected', reviewed_by=?, reviewed_at=?
               WHERE batch_id=? AND status='pending'""",
            (g.user["id"], timestamp, lookup_request["batch_id"]),
        )
        audit(
            db, g.user["display_name"], g.user["id"], "reject", "lookup_request",
            request_id,
            after={"field": lookup_request["field"], "value": lookup_request["value"]},
        )
        db.commit()
        return jsonify(ok=True, batch_status="rejected")

    @app.post("/api/admin/users")
    @admin_required
    def create_user():
        payload = request.get_json(silent=True) or {}
        username = str(payload.get("username", "")).strip()
        display_name = str(payload.get("display_name", "")).strip()
        password = str(payload.get("password", ""))
        role = str(payload.get("role", "staff"))
        if len(username) < 3 or not display_name or len(password) < 12 or role not in {"admin", "staff"}:
            return jsonify(error="Enter a username (at least 3 characters), display name, and password (at least 12 characters)."), 400
        db = get_db()
        try:
            cursor = db.execute(
                """INSERT INTO users (username, display_name, password_hash, role, created_at)
                   VALUES (?, ?, ?, ?, ?)""",
                (username, display_name, generate_password_hash(password), role, now_iso()),
            )
        except sqlite3.IntegrityError:
            return jsonify(error="That username is already in use."), 409
        audit(
            db, g.user["display_name"], g.user["id"], "create", "user", cursor.lastrowid,
            after={"username": username, "display_name": display_name, "role": role},
        )
        db.commit()
        return jsonify(id=cursor.lastrowid), 201

    @app.patch("/api/admin/users/<int:user_id>")
    @admin_required
    def update_user(user_id: int):
        payload = request.get_json(silent=True) or {}
        db = get_db()
        user = db.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
        if not user:
            return jsonify(error="User not found."), 404
        role = str(payload.get("role", user["role"]))
        active = int(bool(payload.get("active", user["active"])))
        if role not in {"admin", "staff"}:
            return jsonify(error="Role must be admin or staff."), 400
        if user_id == g.user["id"] and (not active or role != "admin"):
            return jsonify(error="You cannot deactivate or demote your own account."), 400
        if user["role"] == "admin" and user["active"] and (not active or role != "admin"):
            other_admins = db.execute(
                "SELECT COUNT(*) FROM users WHERE role = 'admin' AND active = 1 AND id != ?",
                (user_id,),
            ).fetchone()[0]
            if not other_admins:
                return jsonify(error="At least one active administrator must remain."), 400
        password = str(payload.get("password", ""))
        if password and len(password) < 12:
            return jsonify(error="Passwords must be at least 12 characters."), 400
        before = {"username": user["username"], "display_name": user["display_name"], "role": user["role"], "active": user["active"]}
        if password:
            db.execute(
                "UPDATE users SET role = ?, active = ?, password_hash = ? WHERE id = ?",
                (role, active, generate_password_hash(password), user_id),
            )
        else:
            db.execute("UPDATE users SET role = ?, active = ? WHERE id = ?", (role, active, user_id))
        after = {**before, "role": role, "active": active, "password_reset": bool(password)}
        audit(db, g.user["display_name"], g.user["id"], "update", "user", user_id, before, after)
        db.commit()
        return jsonify(ok=True)

    @app.get("/api/admin/audit")
    @admin_required
    def admin_audit():
        limit = min(500, max(1, request.args.get("limit", default=200, type=int) or 200))
        rows = get_db().execute(
            "SELECT * FROM audit_log ORDER BY id DESC LIMIT ?", (limit,)
        ).fetchall()
        return jsonify(items=[dict(row) for row in rows])

    return app


def reference_data(workbook_path: Path) -> dict[str, str]:
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    province_areas: dict[str, str] = {}
    if "Locations" in workbook.sheetnames:
        sheet = workbook["Locations"]
        for column in range(1, sheet.max_column + 1):
            area = clean_value(sheet.cell(1, column).value)
            if not area:
                continue
            for row in range(2, sheet.max_row + 1):
                province = clean_value(sheet.cell(row, column).value)
                if province:
                    current = province_areas.get(province)
                    province_areas[province] = area if current in (None, area) else ""
    workbook.close()
    return province_areas


def seed_location_reference(database: sqlite3.Connection, reference_path: Path) -> None:
    if not reference_path.is_file():
        return
    province_areas = reference_data(reference_path)
    for area in sorted(set(province_areas.values()) - {""}):
        add_lookup(database, "area", area, source="location-reference")
    for province, area in province_areas.items():
        if area:
            add_lookup(database, "province", province, source="location-reference", parent_value=area)


def excel_date(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, (int, float)):
        try:
            from openpyxl.utils.datetime import from_excel

            converted = from_excel(value)
            if isinstance(converted, (datetime, date)):
                return converted.date().isoformat() if isinstance(converted, datetime) else converted.isoformat()
        except (ValueError, TypeError, OverflowError):
            return ""
    text = str(value).strip()
    if text.upper() in {"-", "N/A", "NA", "NONE"}:
        return ""
    for date_format in (
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d/%m/%y",
        "%A %d %B %Y",
        "%A %B %d,%Y",
        "%A %B %d, %Y",
        "%d %B %Y",
        "%d-%m-%Y",
    ):
        try:
            return datetime.strptime(text, date_format).date().isoformat()
        except ValueError:
            continue
    return ""


def discover_inventory_sheets(workbook) -> dict[str, tuple[int, dict[str, Any]]]:
    aliases = {
        "description": {"DESCRIPTION"},
        "model": {"MODEL NUMBER", "MODEL"},
        "equipment_brand": {"FREEZER BRAND", "CHILLER BRAND", "MANUFACTURER"},
        "serial_number": {"SERIAL NUMBER", "SERIAL"},
        "store_name": {"STORE NAME", "BUSINESS NAME"},
        "source_location": {"STORE LOCATION", "LOCATION"},
        "business_address": {"BUSINESS ADDRESS"},
        "area": {"AREA"},
        "province": {"PROVINCE"},
        "manager": {"STORE MANAGER", "MANAGER", "CONTACT NAME"},
        "contact": {"STORE CONTACT", "CONTACT NUMBER", "CONTACT"},
        "delivery_date": {"DATE OF DELIVERY", "DATE PLACED", "DELIVERY DATE"},
        "sales_rep": {"SALES REP", "SALESREP/DISTRIBUTOR"},
        "product_line": {"BRANDING"},
    }
    discovered: dict[str, tuple[int, dict[str, int]]] = {}
    for sheet in workbook.worksheets:
        title = clean_value(sheet.title)
        group_headings = {
            clean_value(value)
            for row in sheet.iter_rows(min_row=1, max_row=2, values_only=True)
            for value in row
            if value is not None
        }
        customer_freezer_layout = {
            "CUSTOMER DATA",
            "FREEZER DATA",
            "MISCELLANEOUS",
        }.issubset(group_headings)
        for row_number, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
            headers = {clean_value(value): index for index, value in enumerate(cells) if value is not None}
            if not any("SERIAL NUMBER" in heading for heading in headers):
                continue
            if customer_freezer_layout:
                location_index = headers.get("STORE LOCATION")
                named_area_index = headers.get("AREA")
                region_index = next(
                    (
                        index
                        for index in range(location_index + 1, named_area_index)
                        if cells[index] is None
                    ),
                    None,
                ) if location_index is not None and named_area_index is not None else None
                columns = {
                    "description": headers.get("DESCRIPTION"),
                    "model": headers.get("MODEL NUMBER"),
                    "product_line": headers.get("FREEZER BRAND"),
                    "equipment_brand": headers.get("MANUFACTURER"),
                    "serial_number": headers.get("SERIAL NUMBER"),
                    "store_name": headers.get("STORE NAME"),
                    "business_address": location_index,
                    "area": region_index,
                    "province": named_area_index,
                    "manager": headers.get("STORE MANAGER"),
                    "contact": headers.get("STORE CONTACT"),
                    "sales_rep": headers.get("SALES REP"),
                    "delivery_date": headers.get("DATE OF DELIVERY"),
                    "_customer_freezer_layout": True,
                }
            else:
                columns = {
                    field: next((index for heading, index in headers.items() if heading in options), None)
                    for field, options in aliases.items()
                }
            if columns["serial_number"] is not None and columns["description"] is not None:
                discovered[title] = (
                    row_number,
                    {key: value for key, value in columns.items() if value is not None},
                )
                break
    return discovered


def workbook_records(workbook_bytes: bytes, selected_sheets: list[str], db: sqlite3.Connection) -> tuple[list[str], list[dict[str, Any]]]:
    try:
        from io import BytesIO
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Install project dependencies first with: python -m pip install -r requirements.txt") from exc

    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
    discovered = discover_inventory_sheets(workbook)
    if not discovered:
        workbook.close()
        raise ValueError("No inventory sheets with a description and serial-number column were found.")
    if not selected_sheets:
        workbook.close()
        return list(discovered), []
    unknown = sorted(set(selected_sheets) - set(discovered))
    if unknown:
        workbook.close()
        raise ValueError(f"Selected sheet(s) were not found or are not inventory sheets: {', '.join(unknown)}.")

    province_rows = db.execute(
        "SELECT value, parent_value FROM lookup_values WHERE field = 'province' AND active = 1"
    ).fetchall()
    province_areas = {row["value"]: row["parent_value"] for row in province_rows}
    records: list[dict[str, Any]] = []

    for sheet_name in selected_sheets:
        sheet = next(sheet for sheet in workbook.worksheets if clean_value(sheet.title) == sheet_name)
        header_row, columns = discovered[sheet_name]

        def value_at(row: tuple[Any, ...], field: str) -> Any:
            index = columns.get(field)
            return row[index] if index is not None and index < len(row) else None

        customer_freezer_layout = columns.get("_customer_freezer_layout", False)
        for source_row, row in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            description = clean_value(value_at(row, "description"))
            model = clean_value(value_at(row, "model"))
            brand = clean_value(value_at(row, "equipment_brand"))
            serial = clean_value(value_at(row, "serial_number"))
            store = clean_value(value_at(row, "store_name"))
            source_location = clean_value(value_at(row, "source_location"))
            explicit_address = " ".join(str(value_at(row, "business_address") or "").split())
            explicit_area = clean_value(value_at(row, "area"))
            explicit_province = clean_value(value_at(row, "province"))
            if not any((description, model, brand, serial, store, source_location, explicit_address)):
                continue
            raw_line = clean_value(value_at(row, "product_line"))
            product_line = raw_line if customer_freezer_layout or (sheet_name == "EXPORT" and raw_line) else sheet_name
            raw_location = source_location or explicit_province
            matched_area = explicit_area
            matched_province = explicit_province
            location_needs_review = bool(
                not customer_freezer_layout and raw_location and raw_location not in province_areas
            )
            if not customer_freezer_layout:
                if explicit_province in province_areas:
                    matched_area = province_areas[explicit_province]
                    location_needs_review = False
                elif raw_location and raw_location in province_areas:
                    matched_province = raw_location
                    matched_area = province_areas[raw_location]
                    location_needs_review = False
            delivery_value = value_at(row, "delivery_date")
            raw_date = str(delivery_value or "").strip()
            delivery_date = excel_date(delivery_value)
            status = "in_transit" if sheet_name == "SHIPMENT" else (
                "deployed" if store or source_location or explicit_address else "unassigned"
            )
            records.append(
                {
                    "product_line": product_line,
                    "source_sheet": sheet_name,
                    "source_row": source_row,
                    "customer_freezer_layout": customer_freezer_layout,
                    "equipment_type": description,
                    "equipment_kind": "chiller" if sheet_name == "GUILTLESS" or "CHILLER" in description else "freezer",
                    "equipment_brand": brand,
                    "model": model,
                    "serial_number": serial,
                    "store_name": store,
                    "business_address": explicit_address,
                    "area": matched_area,
                    "province": matched_province,
                    "source_location": raw_location,
                    "location_needs_review": location_needs_review,
                    "location_review_required": bool(
                        customer_freezer_layout and (not explicit_area or not explicit_province)
                    ),
                    "manager": clean_value(value_at(row, "manager")),
                    "contact": " ".join(str(value_at(row, "contact") or "").split()),
                    "delivery_date": delivery_date,
                    "unparsed_delivery_date": raw_date if raw_date and raw_date.upper() not in {"-", "N/A", "NA", "NONE"} and not delivery_date else "",
                    "sales_rep": clean_value(value_at(row, "sales_rep")),
                    "status": status,
                    "notes": "",
                }
            )
    workbook.close()
    return list(discovered), records


IMPORT_LOOKUP_FIELDS = (
    "product_line",
    "equipment_type",
    "equipment_brand",
    "model",
    "area",
    "province",
)


def import_review(db: sqlite3.Connection, records: list[dict[str, Any]]) -> dict[str, Any]:
    values: dict[str, set[str]] = {field: set() for field in IMPORT_LOOKUP_FIELDS}
    counts: dict[tuple[str, str], int] = {}
    for record in records:
        for field in IMPORT_LOOKUP_FIELDS:
            value = clean_value(record.get(field))
            if value:
                values[field].add(value)
                counts[(field, value)] = counts.get((field, value), 0) + 1
    unknown: list[dict[str, Any]] = []
    for field, field_values in values.items():
        approved = {
            row["value"]
            for row in db.execute(
                "SELECT value FROM lookup_values WHERE field = ? AND active = 1", (field,)
            ).fetchall()
        }
        for value in sorted(field_values - approved):
            candidates = difflib.get_close_matches(value, sorted(approved), n=3, cutoff=0.55)
            unknown.append(
                {
                    "field": field,
                    "value": value,
                    "count": counts[(field, value)],
                    "suggestions": candidates,
                }
            )
    location_questions = sorted(
        {
            record["source_location"]
            for record in records
            if record.get("location_needs_review") and record.get("source_location")
        }
    )
    date_warnings = [
        {
            "sheet": record["source_sheet"],
            "row": record["source_row"],
            "value": record["unparsed_delivery_date"],
        }
        for record in records
        if record.get("unparsed_delivery_date")
    ]
    location_warnings = [
        {
            "sheet": record["source_sheet"],
            "row": record["source_row"],
            "store_name": record.get("store_name", ""),
            "business_address": record.get("business_address", ""),
            "missing_fields": [
                field for field in ("area", "province") if not record.get(field)
            ],
        }
        for record in records
        if record.get("location_review_required")
    ]
    return {
        "unknown_values": unknown,
        "location_questions": location_questions,
        "date_warnings": date_warnings,
        "location_warnings": location_warnings,
    }


def import_preview_rows(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    fields = (
        "store_name",
        "business_address",
        "area",
        "province",
        "manager",
        "contact",
        "sales_rep",
        "equipment_type",
        "product_line",
        "equipment_brand",
        "model",
        "serial_number",
        "delivery_date",
    )
    return [
        {
            "source_sheet": record["source_sheet"],
            "source_row": record["source_row"],
            "source_location": record.get("source_location", ""),
            "source_date": record.get("unparsed_delivery_date", ""),
            "location_needs_review": bool(record.get("location_needs_review")),
            "values": {field: record.get(field, "") for field in fields},
        }
        for record in records
    ]


def apply_import_row_edits(
    db: sqlite3.Connection,
    records: list[dict[str, Any]],
    submitted_rows: Any,
    acknowledge_location_warnings: bool,
) -> list[tuple[str, str]]:
    if not isinstance(submitted_rows, list) or len(submitted_rows) != len(records):
        raise ValueError("The reviewed rows do not match this import preview. Preview the workbook again.")

    lookup_fields = ("product_line", "equipment_type", "equipment_brand", "model")
    text_fields = ("store_name", "business_address", "manager", "contact", "sales_rep", "serial_number")
    approved: dict[str, set[str]] = {}
    for field in (*lookup_fields, "area", "province"):
        approved[field] = {
            row["value"]
            for row in db.execute(
                "SELECT value FROM lookup_values WHERE field=? AND active=1", (field,)
            ).fetchall()
        }
    province_areas = {
        row["value"]: row["parent_value"]
        for row in db.execute(
            "SELECT value, parent_value FROM lookup_values WHERE field='province' AND active=1"
        ).fetchall()
    }
    pending: list[tuple[str, str]] = []
    incomplete_locations: list[str] = []

    for record, submitted in zip(records, submitted_rows):
        if not isinstance(submitted, dict):
            raise ValueError("Each reviewed import row must be an object.")
        if (
            submitted.get("source_sheet") != record["source_sheet"]
            or submitted.get("source_row") != record["source_row"]
        ):
            raise ValueError("The reviewed rows do not match this import preview. Preview the workbook again.")
        values = submitted.get("values")
        new_values = submitted.get("new_values", {})
        if not isinstance(values, dict) or not isinstance(new_values, dict):
            raise ValueError(f"Invalid field values for {record['source_sheet']} row {record['source_row']}.")
        row_label = f"{record['source_sheet']} row {record['source_row']}"

        for field in text_fields:
            value = clean_value(values.get(field, record.get(field, "")))
            record[field] = " ".join(value.split()) if field == "business_address" else value

        for field in lookup_fields:
            original_value = clean_value(record.get(field))
            selected = values.get(field, record.get(field, ""))
            if not isinstance(selected, str):
                raise ValueError(f"Choose a valid {field.replace('_', ' ')} for {row_label}.")
            if selected == "__new__":
                value = clean_value(new_values.get(field))
                if not value:
                    raise ValueError(f"Enter the new {field.replace('_', ' ')} for {row_label}.")
                record[field] = value
                pending.append((field, value))
            elif selected == "":
                if original_value:
                    raise ValueError(
                        f"Choose a match or request a new {field.replace('_', ' ')} for {row_label}."
                    )
                record[field] = ""
            elif selected in approved[field]:
                record[field] = selected
            else:
                raise ValueError(f"“{selected}” is not an approved {field.replace('_', ' ')} value.")

        for field in ("area", "province"):
            selected = values.get(field, record.get(field, ""))
            if not isinstance(selected, str) or (selected and selected not in approved[field]):
                raise ValueError(f"Choose an approved {field} for {row_label}.")
            record[field] = selected
        if record["province"] and province_areas.get(record["province"]) != record["area"]:
            raise ValueError(f"Choose a matching Area and Province for {row_label}.")

        raw_date = record.get("unparsed_delivery_date", "")
        date_value = values.get("delivery_date", record.get("delivery_date", ""))
        leave_date_blank = submitted.get("leave_delivery_date_blank") is True
        if date_value:
            try:
                record["delivery_date"] = parse_iso_date(date_value)
            except ValueError as exc:
                raise ValueError(f"Use YYYY-MM-DD for the delivery date in {row_label}.") from exc
            if leave_date_blank:
                raise ValueError(f"Enter a delivery date or leave it blank for {row_label}, not both.")
        elif raw_date and not leave_date_blank:
            raise ValueError(f"Enter a valid delivery date or explicitly leave it blank for {row_label}.")
        else:
            record["delivery_date"] = ""

        if not record["area"] or not record["province"]:
            incomplete_locations.append(row_label)
        if record.get("location_needs_review") and not (
            record["business_address"] or (record["area"] and record["province"])
        ):
            raise ValueError(
                f"Review source location “{record.get('source_location', '')}” in {row_label}: "
                "enter it as a business address or choose an Area and Province."
            )

        record["equipment_kind"] = (
            "chiller"
            if record["source_sheet"] == "GUILTLESS" or "CHILLER" in record["equipment_type"]
            else "freezer"
        )
        record.pop("source_location", None)
        record.pop("location_needs_review", None)
        record.pop("location_review_required", None)
        record.pop("unparsed_delivery_date", None)

    if incomplete_locations and not acknowledge_location_warnings:
        preview = ", ".join(incomplete_locations[:5])
        suffix = " and more" if len(incomplete_locations) > 5 else ""
        raise ValueError(
            f"Review and acknowledge rows with a missing Area or Province before importing: {preview}{suffix}."
        )
    return pending


def ensure_import_store(db: sqlite3.Connection, record: dict[str, Any]) -> int:
    if not record.get("customer_freezer_layout"):
        return ensure_store(db, record)
    exact = db.execute(
        """SELECT id FROM stores
           WHERE store_name=? AND business_address=? AND area=? AND province=?""",
        store_identity(record),
    ).fetchone()
    if exact:
        db.execute(
            """UPDATE stores SET manager=COALESCE(NULLIF(?, ''), manager),
               contact=COALESCE(NULLIF(?, ''), contact),
               sales_rep=COALESCE(NULLIF(?, ''), sales_rep),
               updated_at=? WHERE id=?""",
            (
                record.get("manager", ""),
                record.get("contact", ""),
                record.get("sales_rep", ""),
                now_iso(),
                exact["id"],
            ),
        )
        return exact["id"]
    matches = db.execute(
        """SELECT * FROM stores WHERE store_name=? AND business_address=? ORDER BY id""",
        (clean_value(record.get("store_name")), " ".join(record.get("business_address", "").split()).upper()),
    ).fetchall()
    if len(matches) != 1:
        return ensure_store(db, record)
    existing = matches[0]
    area = record.get("area", "") or existing["area"]
    province = record.get("province", "") or existing["province"]
    manager = record.get("manager", "") or existing["manager"]
    contact = record.get("contact", "") or existing["contact"]
    sales_rep = record.get("sales_rep", "") or existing["sales_rep"]
    db.execute(
        """UPDATE stores SET area=?, province=?, manager=?, contact=?, sales_rep=?, updated_at=? WHERE id=?""",
        (area, province, manager, contact, sales_rep, now_iso(), existing["id"]),
    )
    return existing["id"]


def apply_import_batch(db: sqlite3.Connection, batch: sqlite3.Row, approved_by: sqlite3.Row) -> int:
    records = json.loads(batch["records_json"])
    timestamp = now_iso()
    imported = 0
    for record in records:
        store_id = (
            ensure_import_store(db, record)
            if record.get("store_name")
            else None
        )
        existing = db.execute(
            "SELECT id, store_id, business_address, area, province FROM assets WHERE source_sheet = ? AND source_row = ?",
            (record["source_sheet"], record["source_row"]),
        ).fetchone()
        columns = (
            "store_id", "product_line", "source_sheet", "source_row", "equipment_type", "equipment_kind",
            "equipment_brand", "model", "serial_number", "store_name", "business_address", "area",
            "province", "manager", "contact", "delivery_date", "sales_rep", "status", "notes",
        )
        values = (store_id, *(record.get(field, "") for field in columns[1:]))
        if existing:
            assignments = ", ".join(
                f"{field}=?" for field in columns if field not in {"source_sheet", "source_row"}
            )
            update_values = (store_id, *(record.get(field, "") for field in columns[1:] if field not in {"source_sheet", "source_row"}))
            db.execute(
                f"UPDATE assets SET {assignments}, updated_at=? WHERE id=?",
                (*update_values, timestamp, existing["id"]),
            )
            asset_id = existing["id"]
            changed_location = any(
                existing[field] != record[field]
                for field in ("business_address", "area", "province")
            )
        else:
            placeholders = ", ".join("?" for _ in range(len(columns) + 2))
            cursor = db.execute(
                f"INSERT INTO assets ({', '.join(columns)}, created_at, updated_at) VALUES ({placeholders})",
                (*values, timestamp, timestamp),
            )
            asset_id = cursor.lastrowid
            changed_location = True
        for field in IMPORT_LOOKUP_FIELDS:
            add_lookup(db, field, record.get(field), source="approved-import", created_by=approved_by["id"])
        if changed_location and (
            record.get("business_address") or record.get("area") or record.get("province")
        ):
            db.execute(
                """INSERT INTO location_events
                   (asset_id, business_address, area, province, effective_date, notes, created_by, created_at)
                   VALUES (?, ?, ?, ?, NULL, ?, ?, ?)""",
                (
                    asset_id, record.get("business_address", ""), record.get("area", ""),
                    record.get("province", ""),
                    "Initial placement from reviewed workbook import; source does not provide placement history.",
                    approved_by["id"], timestamp,
                ),
            )
        imported += 1
    db.execute(
        "UPDATE import_batches SET status='complete', completed_at=? WHERE id=?",
        (timestamp, batch["id"]),
    )
    audit(
        db,
        approved_by["display_name"],
        approved_by["id"],
        "approve_import",
        "import_batch",
        batch["id"],
        after={"file_name": batch["file_name"], "records": imported, "sheets": json.loads(batch["sheets_json"])},
    )
    return imported


app = create_app()


def cli() -> None:
    parser = argparse.ArgumentParser(description="Local freezer and chiller inventory server")
    subparsers = parser.add_subparsers(dest="command", required=True)
    subparsers.add_parser("create-admin", help="Create the first administrator account")
    serve_parser = subparsers.add_parser("serve", help="Run the local office web server")
    serve_parser.add_argument("--host", default="0.0.0.0")
    serve_parser.add_argument("--port", type=int, default=8000)
    args = parser.parse_args()

    with app.app_context():
        db = app.get_db()
        if args.command == "create-admin":
            if db.execute("SELECT 1 FROM users WHERE role = 'admin' LIMIT 1").fetchone():
                raise SystemExit("An administrator already exists. Ask an existing admin to create accounts.")
            username = input("Admin username: ").strip()
            display_name = input("Admin display name: ").strip()
            password = getpass.getpass("Password (at least 12 characters): ")
            confirmation = getpass.getpass("Confirm password: ")
            if len(username) < 3 or not display_name or len(password) < 12 or password != confirmation:
                raise SystemExit("Invalid input. Use a username of at least 3 characters, a display name, and matching passwords of at least 12 characters.")
            try:
                db.execute(
                    "INSERT INTO users (username, display_name, password_hash, role, created_at) VALUES (?, ?, ?, 'admin', ?)",
                    (username, display_name, generate_password_hash(password), now_iso()),
                )
                audit(db, display_name, None, "create", "user", username, after={"role": "admin"})
                db.commit()
            except sqlite3.IntegrityError as exc:
                raise SystemExit("That username is already in use.") from exc
            print("Administrator account created.")
        elif args.command == "serve":
            from waitress import serve

            start_backup_scheduler(app)
            print(f"Serving on http://{args.host}:{args.port}")
            serve(app, host=args.host, port=args.port)


if __name__ == "__main__":
    cli()
