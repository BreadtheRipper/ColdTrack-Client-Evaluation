import json
import sqlite3
import tempfile
import unittest
from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from werkzeug.security import generate_password_hash

from app import create_app, run_scheduled_backup


class SorterApiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.database = str(Path(self.temp.name) / "test.sqlite3")
        self.app = create_app(
            {
                "TESTING": True,
                "SECRET_KEY": "test-only-secret",
                "DATABASE": self.database,
                "BACKUP_DIR": str(Path(self.temp.name) / "backups"),
            }
        )
        with self.app.app_context():
            db = self.app.get_db()
            db.execute(
                """INSERT INTO users (username, display_name, password_hash, role, created_at)
                   VALUES (?, ?, ?, 'admin', '2026-01-01T00:00:00')""",
                ("admin", "Admin User", generate_password_hash("correct horse battery staple")),
            )
            db.execute(
                """INSERT INTO users (username, display_name, password_hash, role, created_at)
                   VALUES (?, ?, ?, 'staff', '2026-01-01T00:00:00')""",
                ("staff", "Staff User", generate_password_hash("staff password example")),
            )
            self.area, self.province = db.execute(
                """SELECT a.value AS area_value, p.value AS province_value FROM lookup_values a
                   JOIN lookup_values p ON p.field='province' AND p.parent_value=a.value
                   WHERE a.field='area' ORDER BY a.value, p.value LIMIT 1"""
            ).fetchone()
            self.other_location = db.execute(
                """SELECT a.value AS area_value, p.value AS province_value FROM lookup_values a
                   JOIN lookup_values p ON p.field='province' AND p.parent_value=a.value
                   WHERE a.field='area' AND a.value != ? ORDER BY a.value, p.value LIMIT 1""",
                (self.area,),
            ).fetchone()
            db.commit()
        self.client = self.app.test_client()
        self.csrf = self.client.get("/api/me").json["csrf_token"]
        self.login()

    def tearDown(self):
        self.temp.cleanup()

    def login(self, username="admin", password="correct horse battery staple"):
        response = self.client.post(
            "/api/login",
            json={"username": username, "password": password},
        )
        self.assertEqual(response.status_code, 200, response.json)
        self.csrf = response.json["csrf_token"]

    def request(self, method, path, **kwargs):
        headers = kwargs.pop("headers", {})
        headers["X-CSRF-Token"] = self.csrf
        return getattr(self.client, method)(path, headers=headers, **kwargs)

    def add_lookup(self, field, value):
        response = self.request(
            "post",
            "/api/admin/lookups",
            json={"field": field, "value": value},
        )
        self.assertEqual(response.status_code, 201, response.json)

    def seed_asset_options(self):
        for field, value in (
            ("product_line", "FLAVORITE"),
            ("equipment_type", "2 FT CHEST FREEZER"),
            ("equipment_brand", "RIGID"),
            ("model", "SD220Q"),
            ("store_name", "TEST STORE"),
            ("manager", "ALEX"),
            ("sales_rep", "JAMIE"),
            ("technician", "MAYA"),
        ):
            self.add_lookup(field, value)

    def asset_payload(self, serial="SERIAL-01", area=None, province=None):
        return {
            "product_line": "FLAVORITE",
            "equipment_type": "2 FT CHEST FREEZER",
            "equipment_brand": "RIGID",
            "model": "SD220Q",
            "serial_number": serial,
            "store_name": "TEST STORE",
            "business_address": "10 TEST ROAD",
            "area": area or self.area,
            "province": province or self.province,
            "manager": "ALEX",
            "contact": "555-0100",
            "delivery_date": "2026-01-15",
            "sales_rep": "JAMIE",
            "status": "deployed",
            "notes": "",
        }

    def make_workbook(self, location, delivery_date="2026-01-15"):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "SUNDAYS"
        sheet.append(
            [
                "Description",
                "Model Number",
                "Freezer Brand",
                "Serial Number",
                "Store Name",
                "Store Location",
                "Store Manager",
                "Store Contact",
                "Date of Delivery",
                "Sales Rep",
            ]
        )
        sheet.append(
            [
                "Chest Freezer",
                "MODEL-1",
                "BRAND-1",
                "SERIAL-IMPORT-1",
                "STORE-1",
                location,
                "MANAGER-1",
                "555-0101",
                delivery_date,
                "REP-1",
            ]
        )
        for title in ("EXPORT", "SHIPMENT"):
            workbook.create_sheet(title)
        content = BytesIO()
        workbook.save(content)
        workbook.close()
        return content.getvalue()

    def make_customer_freezer_workbook(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "KLTEST1"
        sheet.merge_cells("A1:F1")
        sheet.merge_cells("H1:K1")
        sheet.merge_cells("M1:N1")
        sheet["A1"] = "CUSTOMER DATA"
        sheet["H1"] = "FREEZER DATA"
        sheet["M1"] = "MISCELLANEOUS"
        sheet.append(
            [
                "Store Name",
                "Store Location",
                None,
                "Area",
                "Store Manager",
                "Store Contact",
                None,
                "Description",
                "Model Number",
                "Serial Number",
                "Freezer Brand",
                None,
                "Sales Rep",
                "Date of Delivery",
            ]
        )
        sheet.append(
            [
                "NEW TEST STORE",
                "10 TEST ROAD",
                "EAST",
                "CUREPE",
                "STORE MANAGER",
                "555-1212",
                None,
                "2FT CHEST",
                "FOCUS 73",
                "NEW-FORMAT-01",
                "ELCOLD",
                None,
                "JANE SMITH",
                datetime(2025, 5, 2),
            ]
        )
        sheet.append(
            [
                "NEW TEST STORE",
                "10 TEST ROAD",
                None,
                None,
                "STORE MANAGER",
                "555-1212",
                None,
                "STANDING FREEZER",
                "VISION 390",
                "NEW-FORMAT-02",
                "ELCOLD",
                None,
                None,
                None,
            ]
        )
        content = BytesIO()
        workbook.save(content)
        workbook.close()
        return content.getvalue()

    def upload_workbook(self, content, sheets=None):
        form = {"workbook": (BytesIO(content), "master.xlsx")}
        if sheets:
            form["sheets"] = sheets
        return self.request(
            "post",
            "/api/imports/preview",
            data=form,
            content_type="multipart/form-data",
        )

    def import_choices(self, review):
        choices = {}
        for item in review["unknown_values"]:
            choices.setdefault(item["field"], {})[item["value"]] = "__new__"
        return choices

    def test_reference_lists_are_seeded_from_locations_only(self):
        lookups = self.client.get("/api/lookups").json
        self.assertIn(self.area, [item["value"] for item in lookups["area"]])
        self.assertTrue(
            any(
                item["value"] == self.province and item["parent_value"] == self.area
                for item in lookups["province"]
            )
        )
        with self.app.app_context():
            db = self.app.get_db()
            self.assertEqual(db.execute("SELECT COUNT(*) FROM assets").fetchone()[0], 0)
            self.assertEqual(
                db.execute(
                    "SELECT COUNT(*) FROM lookup_values WHERE field NOT IN ('area', 'province')"
                ).fetchone()[0],
                0,
            )

    def test_grouped_inventory_keeps_serial_level_records(self):
        self.seed_asset_options()
        first = self.request("post", "/api/assets", json=self.asset_payload("SERIAL-01"))
        second = self.request("post", "/api/assets", json=self.asset_payload("SERIAL-02"))
        self.assertEqual(first.status_code, 201, first.json)
        self.assertEqual(second.status_code, 201, second.json)

        groups = self.client.get("/api/assets?grouped=1").json
        self.assertEqual(groups["total"], 1)
        self.assertEqual(groups["total_units"], 2)
        self.assertEqual(groups["items"][0]["unit_count"], 2)
        address_filter = self.client.get("/api/assets?grouped=1&business_address=TEST")
        self.assertEqual(address_filter.json["total_units"], 2)
        members = self.client.get(
            f"/api/assets?ids={groups['items'][0]['member_ids']}&per_page=500"
        ).json
        self.assertEqual({item["serial_number"] for item in members["items"]}, {"SERIAL-01", "SERIAL-02"})

    def test_equipment_can_be_added_to_stock_or_store_without_service(self):
        self.seed_asset_options()
        stock_payload = self.asset_payload(serial="STOCK-01")
        stock_payload.update(
            store_name="",
            business_address="",
            area="",
            province="",
            manager="",
            contact="",
            sales_rep="",
            status="unassigned",
            delivery_date="",
            notes="Stored in warehouse",
        )
        stock = self.request("post", "/api/assets", json=stock_payload)
        self.assertEqual(stock.status_code, 201, stock.json)
        available = self.client.get("/api/assets?available=1&per_page=500")
        self.assertEqual(available.json["total"], 1)
        self.assertEqual(available.json["items"][0]["serial_number"], "STOCK-01")

        store_asset = self.request("post", "/api/assets", json=self.asset_payload("STORE-BASE"))
        self.assertEqual(store_asset.status_code, 201, store_asset.json)
        store_id = self.client.get(f"/api/assets/{store_asset.json['id']}").json["asset"]["store_id"]

        assigned = self.request(
            "post",
            f"/api/stores/{store_id}/freezers",
            json={"asset_id": stock.json["id"]},
        )
        self.assertEqual(assigned.status_code, 201, assigned.json)
        direct = self.request(
            "post",
            f"/api/stores/{store_id}/freezers",
            json={
                "new_freezer": {
                    "product_line": "FLAVORITE",
                    "equipment_type": "2 FT CHEST FREEZER",
                    "equipment_brand": "RIGID",
                    "model": "SD220Q",
                    "serial_number": "STORE-ONLY-01",
                    "delivery_date": "2026-04-05",
                    "notes": "No service performed yet",
                }
            },
        )
        self.assertEqual(direct.status_code, 201, direct.json)

        store_view = self.client.get(f"/api/stores/{store_id}").json
        self.assertEqual(
            {freezer["serial_number"] for freezer in store_view["freezers"]},
            {"STORE-BASE", "STOCK-01", "STORE-ONLY-01"},
        )
        self.assertTrue(all(freezer["service_count"] == 0 for freezer in store_view["freezers"]))
        current = self.client.get("/api/assets?current_only=1&per_page=100").json
        self.assertEqual(current["total"], 3)
        previous = self.client.get(f"/api/assets/{stock.json['id']}").json
        self.assertEqual(previous["asset"]["placement_active"], 0)
        self.assertEqual(previous["services"], [])

    def test_service_and_location_history_filters_and_audit(self):
        self.seed_asset_options()
        created = self.request("post", "/api/assets", json=self.asset_payload())
        self.assertEqual(created.status_code, 201, created.json)
        asset_id = created.json["id"]

        service = self.request(
            "post",
            f"/api/assets/{asset_id}/services",
            json={
                "service_date": "2026-02-02",
                "technician": "MAYA",
                "repairs": "Replaced compressor",
                "parts": "Compressor",
                "notes": "Test visit",
            },
        )
        self.assertEqual(service.status_code, 201, service.json)
        service_id = service.json["id"]
        updated_service = self.request(
            "patch",
            f"/api/services/{service_id}",
            json={
                "service_date": "2026-02-03",
                "technician": "MAYA",
                "repairs": "Replaced compressor and tested",
                "parts": "Compressor",
                "notes": "Corrected visit note",
            },
        )
        self.assertEqual(updated_service.status_code, 200, updated_service.json)

        detail = self.client.get(f"/api/assets/{asset_id}")
        self.assertEqual(len(detail.json["services"]), 1)
        self.assertEqual(detail.json["services"][0]["service_date"], "2026-02-03")
        self.assertEqual(len(detail.json["locations"]), 1)
        filtered = self.client.get(
            f"/api/assets?technician=MAYA&province={self.province}&sort=province"
        )
        self.assertEqual(filtered.json["total"], 1)

        if self.other_location:
            area, province = self.other_location
            changed = self.asset_payload(area=area, province=province)
            changed.update(
                location_effective_date="2026-03-01",
                location_notes="Moved for testing",
            )
            updated = self.request("patch", f"/api/assets/{asset_id}", json=changed)
            self.assertEqual(updated.status_code, 200, updated.json)
            updated_detail = self.client.get(f"/api/assets/{asset_id}").json
            self.assertEqual(len(updated_detail["locations"]), 2)
            self.assertEqual(updated_detail["asset"]["province"], province)

        audit = self.client.get("/api/admin/audit")
        self.assertGreaterEqual(len(audit.json["items"]), 3)

    def test_store_centered_service_entry_and_move_isolates_placement_history(self):
        self.seed_asset_options()
        service = self.request(
            "post",
            "/api/services",
            json={
                "new_store": {
                    "store_name": "BRANCH STORE",
                    "business_address": "10 FIRST ROAD",
                    "area": self.area,
                    "province": self.province,
                    "manager": "MANAGER ONE",
                    "contact": "555-1000",
                },
                "new_freezer": {
                    "product_line": "FLAVORITE",
                    "equipment_type": "2 FT CHEST FREEZER",
                    "equipment_brand": "RIGID",
                    "model": "SD220Q",
                    "serial_number": "MOVE-SERIAL",
                },
                "service_date": "2026-05-05",
                "technician": "MAYA",
                "repairs": "Compressor checked",
                "parts": "Relay",
                "notes": "Working normally",
            },
        )
        self.assertEqual(service.status_code, 201, service.json)
        original_asset_id = service.json["asset_id"]
        original_store_id = service.json["store_id"]

        store_view = self.client.get(f"/api/stores/{original_store_id}").json
        self.assertEqual(len(store_view["freezers"]), 1)
        self.assertEqual(store_view["freezers"][0]["latest_service"], "2026-05-05")
        self.assertIn("Compressor checked", store_view["freezers"][0]["latest_service_details"])

        repeat = self.request(
            "post",
            "/api/services",
            json={
                "store_id": original_store_id,
                "asset_id": original_asset_id,
                "service_date": "2026-06-06",
                "technician": "MAYA",
                "repairs": "Thermostat adjusted",
                "parts": "",
                "notes": "",
            },
        )
        self.assertEqual(repeat.status_code, 201, repeat.json)
        moved = self.request(
            "post",
            f"/api/assets/{original_asset_id}/move",
            json={
                "new_store": {
                    "store_name": "DESTINATION STORE",
                    "business_address": "99 SECOND ROAD",
                    "area": self.area,
                    "province": self.province,
                    "manager": "",
                    "contact": "",
                },
                "effective_date": "2026-07-01",
            },
        )
        self.assertEqual(moved.status_code, 201, moved.json)
        destination_store = self.client.get(f"/api/stores/{moved.json['store_id']}").json
        self.assertEqual(destination_store["store"]["store_name"], "DESTINATION STORE")
        self.assertEqual(len(destination_store["freezers"]), 1)
        self.assertEqual(destination_store["freezers"][0]["serial_number"], "MOVE-SERIAL")
        self.assertEqual(destination_store["freezers"][0]["service_count"], 0)
        old_store_report = self.client.get(f"/api/reports/services?store_id={original_store_id}")
        self.assertEqual(old_store_report.status_code, 200, old_store_report.json)
        self.assertEqual(old_store_report.json["total_services"], 2)
        self.assertTrue(all(row["store_name"] == "BRANCH STORE" for row in old_store_report.json["items"]))
        serial_report = self.client.get("/api/reports/services?serial_number=move-ser")
        self.assertEqual(serial_report.status_code, 200, serial_report.json)
        self.assertEqual(serial_report.json["total_services"], 2)
        self.assertEqual(serial_report.json["items"][0]["asset_id"], original_asset_id)
        destination_report = self.client.get(f"/api/reports/services?store_id={moved.json['store_id']}")
        self.assertEqual(destination_report.json["total_services"], 0)
        invalid_store_report = self.client.get("/api/reports/services?store_id=invalid")
        self.assertEqual(invalid_store_report.status_code, 400)
        destination_service = self.request(
            "post",
            "/api/services",
            json={
                "store_id": moved.json["store_id"],
                "asset_id": moved.json["id"],
                "service_date": "2026-08-01",
                "technician": "MAYA",
                "repairs": "Fan replaced",
                "parts": "Fan motor",
                "notes": "",
            },
        )
        self.assertEqual(destination_service.status_code, 201, destination_service.json)
        all_placement_services = self.client.get("/api/reports/services?serial_number=MOVE-SERIAL")
        self.assertEqual(all_placement_services.json["total_services"], 3)
        previous_store = self.client.get(f"/api/stores/{original_store_id}").json
        self.assertEqual(previous_store["freezers"], [])
        original_history = self.client.get(f"/api/assets/{original_asset_id}").json
        self.assertEqual(len(original_history["services"]), 2)
        moved_unit_report = self.client.get(f"/api/reports/services/{moved.json['id']}")
        self.assertEqual(moved_unit_report.status_code, 200, moved_unit_report.json)
        self.assertEqual(len(moved_unit_report.json["services"]), 3)
        self.assertEqual(
            {service["store_name"] for service in moved_unit_report.json["services"]},
            {"BRANCH STORE", "DESTINATION STORE"},
        )
        with self.app.app_context():
            new_placement = self.app.get_db().execute(
                "SELECT moved_from_asset_id, placement_active FROM assets WHERE id=?",
                (moved.json["id"],),
            ).fetchone()
            self.assertEqual(new_placement["moved_from_asset_id"], original_asset_id)
            self.assertEqual(new_placement["placement_active"], 1)

    def test_store_info_saves_sales_rep_and_notes(self):
        self.seed_asset_options()
        created = self.request("post", "/api/assets", json=self.asset_payload(serial="STORE-INFO-01"))
        self.assertEqual(created.status_code, 201, created.json)
        store = self.client.get("/api/stores").json["items"][0]
        self.assertEqual(store["sales_rep"], "JAMIE")

        updated = self.request(
            "patch",
            f"/api/stores/{store['id']}",
            json={
                "sales_rep": "TAYLOR",
                "notes": "Call the manager before scheduling service.",
            },
        )
        self.assertEqual(updated.status_code, 200, updated.json)
        detail = self.client.get(f"/api/stores/{store['id']}").json["store"]
        self.assertEqual(detail["sales_rep"], "TAYLOR")
        self.assertEqual(detail["notes"], "Call the manager before scheduling service.")
        search = self.client.get("/api/stores?search=scheduling%20service").json["items"]
        self.assertEqual([result["id"] for result in search], [store["id"]])
        with self.app.app_context():
            asset = self.app.get_db().execute(
                "SELECT sales_rep FROM assets WHERE id=?", (created.json["id"],)
            ).fetchone()
            self.assertEqual(asset["sales_rep"], "TAYLOR")

    def test_existing_store_database_migrates_sales_rep_and_notes_columns(self):
        legacy_database = str(Path(self.temp.name) / "legacy.sqlite3")
        with sqlite3.connect(legacy_database) as db:
            db.execute(
                """CREATE TABLE stores (
                    id INTEGER PRIMARY KEY,
                    store_name TEXT NOT NULL,
                    business_address TEXT NOT NULL DEFAULT '',
                    area TEXT NOT NULL DEFAULT '',
                    province TEXT NOT NULL DEFAULT '',
                    manager TEXT NOT NULL DEFAULT '',
                    contact TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                )"""
            )
            db.execute(
                """INSERT INTO stores
                   (store_name, business_address, created_at, updated_at)
                   VALUES ('LEGACY STORE', '1 OLD ROAD', '2025-01-01', '2025-01-01')"""
            )
        db.close()
        migrated_app = create_app(
            {
                "TESTING": True,
                "SECRET_KEY": "migration-test-secret",
                "DATABASE": legacy_database,
                "BACKUP_DIR": str(Path(self.temp.name) / "migration-backups"),
            }
        )
        with migrated_app.app_context():
            store = migrated_app.get_db().execute("SELECT * FROM stores").fetchone()
            self.assertEqual(store["sales_rep"], "")
            self.assertEqual(store["notes"], "")

    def test_admin_can_rename_and_deactivate_dropdown_values(self):
        self.seed_asset_options()
        created = self.request("post", "/api/assets", json=self.asset_payload())
        self.assertEqual(created.status_code, 201, created.json)
        service = self.request(
            "post",
            f"/api/assets/{created.json['id']}/services",
            json={
                "service_date": "2026-02-02",
                "technician": "MAYA",
                "repairs": "",
                "parts": "",
                "notes": "",
            },
        )
        lookups = self.client.get("/api/admin/lookups").json["items"]
        technician_id = next(item["id"] for item in lookups if item["field"] == "technician" and item["value"] == "MAYA")
        renamed = self.request(
            "patch",
            f"/api/admin/lookups/{technician_id}",
            json={"value": "MAYA TECH"},
        )
        self.assertEqual(renamed.status_code, 200, renamed.json)
        history = self.client.get(f"/api/assets/{created.json['id']}").json["services"]
        self.assertEqual(history[0]["technician"], "MAYA TECH")
        deactivated = self.request(
            "patch",
            f"/api/admin/lookups/{technician_id}",
            json={"active": False},
        )
        self.assertEqual(deactivated.status_code, 200, deactivated.json)
        active_technicians = self.client.get("/api/lookups").json["technician"]
        self.assertNotIn("MAYA TECH", [item["value"] for item in active_technicians])
        self.assertEqual(service.status_code, 201, service.json)

    def test_admin_can_rename_freezer_brand_and_manufacturer(self):
        self.seed_asset_options()
        created = self.request("post", "/api/assets", json=self.asset_payload())
        self.assertEqual(created.status_code, 201, created.json)
        lookups = self.client.get("/api/admin/lookups").json["items"]
        for field, old_value, new_value in (
            ("product_line", "FLAVORITE", "NEW FREEZER BRAND"),
            ("equipment_brand", "RIGID", "NEW MANUFACTURER"),
        ):
            lookup_id = next(
                item["id"] for item in lookups
                if item["field"] == field and item["value"] == old_value
            )
            renamed = self.request(
                "patch",
                f"/api/admin/lookups/{lookup_id}",
                json={"value": new_value},
            )
            self.assertEqual(renamed.status_code, 200, renamed.json)

        asset = self.client.get(f"/api/assets/{created.json['id']}").json["asset"]
        self.assertEqual(asset["product_line"], "NEW FREEZER BRAND")
        self.assertEqual(asset["equipment_brand"], "NEW MANUFACTURER")

    def test_admin_can_create_edit_reassign_and_deactivate_area_province_values(self):
        area_response = self.request(
            "post",
            "/api/admin/lookups",
            json={"field": "area", "value": "TEST AREA"},
        )
        self.assertEqual(area_response.status_code, 201, area_response.json)
        missing_parent = self.request(
            "post",
            "/api/admin/lookups",
            json={"field": "province", "value": "TEST PROVINCE"},
        )
        self.assertEqual(missing_parent.status_code, 400)
        province_response = self.request(
            "post",
            "/api/admin/lookups",
            json={"field": "province", "value": "TEST PROVINCE", "parent_value": "TEST AREA"},
        )
        self.assertEqual(province_response.status_code, 201, province_response.json)
        self.seed_asset_options()
        service = self.request(
            "post",
            "/api/services",
            json={
                "new_store": {
                    "store_name": "AREA TEST STORE",
                    "business_address": "1 TEST ROAD",
                    "area": "TEST AREA",
                    "province": "TEST PROVINCE",
                    "manager": "",
                    "contact": "",
                },
                "new_freezer": {
                    "product_line": "FLAVORITE",
                    "equipment_type": "2 FT CHEST FREEZER",
                    "equipment_brand": "RIGID",
                    "model": "SD220Q",
                    "serial_number": "AREA-TEST-1",
                },
                "service_date": "2026-04-01",
                "technician": "MAYA",
            },
        )
        self.assertEqual(service.status_code, 201, service.json)

        lookups = self.client.get("/api/admin/lookups").json["items"]
        area_id = next(item["id"] for item in lookups if item["field"] == "area" and item["value"] == "TEST AREA")
        province_id = next(
            item["id"] for item in lookups if item["field"] == "province" and item["value"] == "TEST PROVINCE"
        )
        renamed_area = self.request(
            "patch",
            f"/api/admin/lookups/{area_id}",
            json={"value": "RENAMED AREA"},
        )
        self.assertEqual(renamed_area.status_code, 200, renamed_area.json)
        other_area = self.other_location["area_value"]
        reassigned = self.request(
            "patch",
            f"/api/admin/lookups/{province_id}",
            json={"value": "TEST PROVINCE", "parent_value": other_area},
        )
        self.assertEqual(reassigned.status_code, 200, reassigned.json)
        renamed_province = self.request(
            "patch",
            f"/api/admin/lookups/{province_id}",
            json={"value": "RENAMED PROVINCE", "parent_value": other_area},
        )
        self.assertEqual(renamed_province.status_code, 200, renamed_province.json)
        store = self.client.get(f"/api/stores/{service.json['store_id']}").json["store"]
        self.assertEqual(store["area"], other_area)
        self.assertEqual(store["province"], "RENAMED PROVINCE")
        with self.app.app_context():
            province = self.app.get_db().execute(
                "SELECT parent_value FROM lookup_values WHERE id=?",
                (province_id,),
            ).fetchone()
            self.assertEqual(province["parent_value"], other_area)

        deactivated_area = self.request(
            "patch",
            f"/api/admin/lookups/{next(item['id'] for item in lookups if item['field'] == 'area' and item['value'] == other_area)}",
            json={"active": False},
        )
        self.assertEqual(deactivated_area.status_code, 200, deactivated_area.json)
        active_lookups = self.client.get("/api/lookups").json
        self.assertNotIn(other_area, [item["value"] for item in active_lookups["area"]])
        self.assertNotIn("RENAMED PROVINCE", [item["value"] for item in active_lookups["province"]])

    def test_admin_can_backup_and_reset_data_while_retaining_current_account(self):
        self.add_lookup("technician", "BACKUP TECH")
        schedule = self.client.get("/api/admin/backups").json
        self.assertEqual(schedule["backup_time"], "02:00")
        invalid_schedule = self.request(
            "patch",
            "/api/admin/backups/schedule",
            json={"time": "25:00"},
        )
        self.assertEqual(invalid_schedule.status_code, 400)
        changed_schedule = self.request(
            "patch",
            "/api/admin/backups/schedule",
            json={"time": "03:15"},
        )
        self.assertEqual(changed_schedule.status_code, 200, changed_schedule.json)
        created = self.request("post", "/api/admin/backups", json={})
        self.assertEqual(created.status_code, 201, created.json)
        downloaded = self.client.get(f"/api/admin/backups/{created.json['filename']}")
        self.assertEqual(downloaded.status_code, 200)
        self.assertTrue(downloaded.data.startswith(b"SQLite format 3"))
        downloaded.close()

        missing_confirmation = self.request("post", "/api/admin/reset", json={})
        self.assertEqual(missing_confirmation.status_code, 400)
        reset = self.request(
            "post",
            "/api/admin/reset",
            json={"confirmation": "RESET"},
        )
        self.assertEqual(reset.status_code, 200, reset.json)
        self.assertTrue((Path(self.temp.name) / "backups" / reset.json["backup"]).is_file())
        backup_db = sqlite3.connect(Path(self.temp.name) / "backups" / reset.json["backup"])
        try:
            self.assertGreater(
                backup_db.execute("SELECT COUNT(*) FROM lookup_values").fetchone()[0],
                0,
            )
            self.assertEqual(backup_db.execute("SELECT COUNT(*) FROM users").fetchone()[0], 2)
        finally:
            backup_db.close()
        with self.app.app_context():
            db = self.app.get_db()
            for table in ("stores", "assets", "service_events", "location_events", "lookup_values", "audit_log"):
                self.assertEqual(db.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0], 0, table)
            users = db.execute("SELECT username FROM users").fetchall()
            self.assertEqual([row["username"] for row in users], ["admin"])
        self.assertEqual(self.client.get("/api/me").json["user"]["username"], "admin")
        restarted_app = create_app(
            {
                "TESTING": True,
                "SECRET_KEY": "test-only-secret",
                "DATABASE": self.database,
                "BACKUP_DIR": str(Path(self.temp.name) / "backups"),
            }
        )
        with restarted_app.app_context():
            self.assertEqual(
                restarted_app.get_db().execute("SELECT COUNT(*) FROM lookup_values").fetchone()[0],
                0,
            )

    def test_daily_backup_runs_once_after_configured_local_time(self):
        scheduled_time = datetime.now().astimezone().replace(
            hour=23,
            minute=59,
            second=0,
            microsecond=0,
        )
        self.assertTrue(run_scheduled_backup(self.app, scheduled_time))
        self.assertFalse(run_scheduled_backup(self.app, scheduled_time))
        backups = list((Path(self.temp.name) / "backups").glob("coldtrack-*.sqlite3"))
        self.assertEqual(len(backups), 1)
        with self.app.app_context():
            last_run = self.app.get_db().execute(
                "SELECT value FROM app_settings WHERE key='backup_last_run_date'"
            ).fetchone()
            self.assertEqual(last_run["value"], scheduled_time.date().isoformat())

    def test_invalid_managed_values_and_area_province_pairs_are_rejected(self):
        invalid = self.request("post", "/api/assets", json=self.asset_payload())
        self.assertEqual(invalid.status_code, 400)
        self.assertIn("managed list", invalid.json["error"])
        self.seed_asset_options()
        if self.other_location:
            payload = self.asset_payload(province=self.other_location["province_value"])
            response = self.request("post", "/api/assets", json=payload)
            self.assertEqual(response.status_code, 400)
            self.assertIn("does not belong", response.json["error"].lower())

    def test_staff_can_edit_assets_but_cannot_administer(self):
        self.seed_asset_options()
        staff_client = self.app.test_client()
        staff_client.get("/api/me")
        login = staff_client.post(
            "/api/login",
            json={"username": "staff", "password": "staff password example"},
        )
        token = login.json["csrf_token"]
        response = staff_client.get("/api/admin/users")
        self.assertEqual(response.status_code, 403)
        change = staff_client.post(
            "/api/assets",
            headers={"X-CSRF-Token": token},
            json=self.asset_payload(serial="STAFF-01"),
        )
        self.assertEqual(change.status_code, 201, change.json)

    def test_import_requires_location_clarification_and_admin_approval(self):
        content = self.make_workbook("12 UNKNOWN ROAD", "Unrecognized date")
        discovery = self.upload_workbook(content)
        self.assertEqual(discovery.status_code, 200, discovery.json)
        self.assertEqual(discovery.json["sheets"], ["SUNDAYS"])

        preview = self.upload_workbook(content, ["SUNDAYS"])
        self.assertEqual(preview.status_code, 201, preview.json)
        batch_id = preview.json["batch_id"]
        self.assertEqual(preview.json["review"]["location_questions"], ["12 UNKNOWN ROAD"])
        self.assertEqual(
            preview.json["review"]["date_warnings"],
            [{"sheet": "SUNDAYS", "row": 2, "value": "Unrecognized date"}],
        )
        choices = self.import_choices(preview.json["review"])
        rejected = self.request(
            "post",
            f"/api/imports/{batch_id}/commit",
            json={"choices": choices, "locations": {}},
        )
        self.assertEqual(rejected.status_code, 400)
        self.assertIn("delivery date", rejected.json["error"].lower())

        unresolved = self.request(
            "post",
            f"/api/imports/{batch_id}/commit",
            json={
                "choices": choices,
                "date_resolutions": {
                    "SUNDAYS:2": {"normalized": "2026-01-15", "leave_blank": False}
                },
            },
        )
        self.assertEqual(unresolved.status_code, 400)
        self.assertIn("clarify", unresolved.json["error"].lower())

        submitted = self.request(
            "post",
            f"/api/imports/{batch_id}/commit",
            json={
                "choices": choices,
                "date_resolutions": {
                    "SUNDAYS:2": {"normalized": "2026-01-15", "leave_blank": False}
                },
                "locations": {
                    "12 UNKNOWN ROAD": {
                        "kind": "address",
                        "address": "12 UNKNOWN ROAD",
                    }
                },
            },
        )
        self.assertEqual(submitted.status_code, 200, submitted.json)
        self.assertEqual(submitted.json["status"], "awaiting_approval")
        with self.app.app_context():
            self.assertEqual(
                self.app.get_db().execute("SELECT COUNT(*) FROM assets").fetchone()[0],
                0,
            )

        requests = self.client.get("/api/admin/import-requests").json["items"]
        self.assertGreater(len(requests), 0)
        for item in requests:
            approved = self.request(
                "post",
                f"/api/admin/import-requests/{item['id']}/approve",
                json={},
            )
            self.assertEqual(approved.status_code, 200, approved.json)
        with self.app.app_context():
            db = self.app.get_db()
            asset = db.execute("SELECT * FROM assets").fetchone()
            batch = db.execute("SELECT status FROM import_batches WHERE id=?", (batch_id,)).fetchone()
            self.assertEqual(batch["status"], "complete")
            self.assertEqual(asset["business_address"], "12 UNKNOWN ROAD")
            self.assertEqual(asset["area"], "")
            self.assertEqual(asset["province"], "")
            self.assertEqual(asset["source_sheet"], "SUNDAYS")
            self.assertEqual(asset["source_row"], 2)
            self.assertEqual(asset["delivery_date"], "2026-01-15")

    def test_import_uses_exact_province_and_rejects_invalid_choice_shapes(self):
        content = self.make_workbook(self.province)
        preview = self.upload_workbook(content, ["SUNDAYS"])
        self.assertEqual(preview.status_code, 201, preview.json)
        self.assertEqual(preview.json["review"]["location_questions"], [])
        malformed = self.request(
            "post",
            f"/api/imports/{preview.json['batch_id']}/commit",
            json={"choices": ["invalid"], "locations": {}},
        )
        self.assertEqual(malformed.status_code, 400)
        self.assertIn("mappings", malformed.json["error"])

    def test_customer_freezer_workbook_import_maps_columns_and_preserves_blank_locations(self):
        content = self.make_customer_freezer_workbook()
        discovery = self.upload_workbook(content)
        self.assertEqual(discovery.status_code, 200, discovery.json)
        self.assertEqual(discovery.json["sheets"], ["KLTEST1"])
        preview = self.upload_workbook(content, ["KLTEST1"])
        self.assertEqual(preview.status_code, 201, preview.json)
        self.assertEqual(preview.json["record_count"], 2)
        review = preview.json["review"]
        self.assertEqual(
            review["location_warnings"],
            [
                {
                    "sheet": "KLTEST1",
                    "row": 4,
                    "store_name": "NEW TEST STORE",
                    "business_address": "10 TEST ROAD",
                    "missing_fields": ["area", "province"],
                }
            ],
        )
        self.assertIn(
            ("area", "EAST"),
            [(item["field"], item["value"]) for item in review["unknown_values"]],
        )
        self.assertFalse(
            any(item["field"] in {"store_name", "manager", "sales_rep"} for item in review["unknown_values"])
        )

        with self.app.app_context():
            batch = self.app.get_db().execute(
                "SELECT records_json FROM import_batches WHERE id=?",
                (preview.json["batch_id"],),
            ).fetchone()
            import_records = json.loads(batch["records_json"])
        self.assertEqual(import_records[0]["product_line"], "ELCOLD")
        self.assertEqual(import_records[0]["equipment_brand"], "")
        self.assertEqual(import_records[0]["business_address"], "10 TEST ROAD")
        self.assertEqual(import_records[0]["area"], "EAST")
        self.assertEqual(import_records[0]["province"], "CUREPE")
        self.assertEqual(import_records[0]["sales_rep"], "JANE SMITH")
        self.assertEqual(import_records[1]["area"], "")
        self.assertEqual(import_records[1]["province"], "")

        with self.app.app_context():
            curepe_area = self.app.get_db().execute(
                "SELECT parent_value FROM lookup_values WHERE field='province' AND value='CUREPE' AND active=1"
            ).fetchone()["parent_value"]
        choices = {}
        for item in review["unknown_values"]:
            choices.setdefault(item["field"], {})[item["value"]] = (
                curepe_area if item["field"] == "area" and item["value"] == "EAST" else "__new__"
            )
        batch_id = preview.json["batch_id"]
        missing_ack = self.request(
            "post",
            f"/api/imports/{batch_id}/commit",
            json={"choices": choices},
        )
        self.assertEqual(missing_ack.status_code, 400)
        self.assertIn("acknowledge", missing_ack.json["error"].lower())
        submitted = self.request(
            "post",
            f"/api/imports/{batch_id}/commit",
            json={"choices": choices, "acknowledge_location_warnings": True},
        )
        self.assertEqual(submitted.status_code, 200, submitted.json)
        self.assertEqual(submitted.json["status"], "awaiting_approval")
        requests = self.client.get("/api/admin/import-requests").json["items"]
        for item in requests:
            approved = self.request(
                "post",
                f"/api/admin/import-requests/{item['id']}/approve",
                json={},
            )
            self.assertEqual(approved.status_code, 200, approved.json)
        with self.app.app_context():
            db = self.app.get_db()
            assets = db.execute(
                "SELECT product_line, equipment_brand, area, province, business_address, sales_rep, store_id FROM assets ORDER BY id"
            ).fetchall()
            self.assertEqual(len(assets), 2)
            self.assertEqual(assets[0]["product_line"], "ELCOLD")
            self.assertEqual(assets[0]["equipment_brand"], "")
            self.assertEqual(assets[0]["area"], curepe_area)
            self.assertEqual(assets[0]["province"], "CUREPE")
            self.assertEqual(assets[1]["area"], "")
            self.assertEqual(assets[1]["province"], "")
            self.assertEqual(assets[0]["business_address"], "10 TEST ROAD")
            self.assertEqual(assets[0]["sales_rep"], "JANE SMITH")
            self.assertEqual(assets[0]["store_id"], assets[1]["store_id"])
            self.assertEqual(db.execute("SELECT COUNT(*) FROM stores").fetchone()[0], 1)
            store = db.execute("SELECT sales_rep FROM stores WHERE id=?", (assets[0]["store_id"],)).fetchone()
            self.assertEqual(store["sales_rep"], "JANE SMITH")

    def test_row_based_import_allows_individual_edits_and_validates_geography(self):
        preview = self.upload_workbook(self.make_customer_freezer_workbook(), ["KLTEST1"])
        self.assertEqual(preview.status_code, 201, preview.json)
        rows = preview.json["rows"]
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["values"]["area"], "EAST")
        self.assertEqual(rows[0]["values"]["province"], "CUREPE")

        lookup_fields = ("product_line", "equipment_type", "equipment_brand", "model")
        for index, row in enumerate(rows, start=1):
            row["values"]["store_name"] = "EDITED FIRST STORE" if index == 1 else "NEW TEST STORE"
            row["values"]["manager"] = f"ROW MANAGER {index}"
            row["values"]["province"] = self.province if index == 1 else ""
            row["values"]["area"] = self.area if index == 1 else ""
            row["new_values"] = {}
            for field in lookup_fields:
                row["values"][field] = "__new__"
                row["new_values"][field] = f"ROW {index} {field.upper()}"

        batch_id = preview.json["batch_id"]
        unresolved = json.loads(json.dumps(rows))
        unresolved[0]["values"]["model"] = ""
        unresolved[0]["new_values"]["model"] = ""
        unresolved_result = self.request(
            "post",
            f"/api/imports/{batch_id}/commit-rows",
            json={"rows": unresolved, "acknowledge_location_warnings": True},
        )
        self.assertEqual(unresolved_result.status_code, 400)
        self.assertIn("request a new model", unresolved_result.json["error"].lower())

        mismatch = json.loads(json.dumps(rows))
        mismatch[0]["values"]["area"] = self.other_location["area_value"]
        rejected = self.request(
            "post",
            f"/api/imports/{batch_id}/commit-rows",
            json={"rows": mismatch, "acknowledge_location_warnings": True},
        )
        self.assertEqual(rejected.status_code, 400)
        self.assertIn("matching area and province", rejected.json["error"].lower())

        missing_ack = self.request(
            "post",
            f"/api/imports/{batch_id}/commit-rows",
            json={"rows": rows, "acknowledge_location_warnings": False},
        )
        self.assertEqual(missing_ack.status_code, 400)
        self.assertIn("acknowledge", missing_ack.json["error"].lower())

        submitted = self.request(
            "post",
            f"/api/imports/{batch_id}/commit-rows",
            json={"rows": rows, "acknowledge_location_warnings": True},
        )
        self.assertEqual(submitted.status_code, 200, submitted.json)
        self.assertEqual(submitted.json["status"], "awaiting_approval")
        requests = self.client.get("/api/admin/import-requests").json["items"]
        self.assertEqual(len(requests), 8)
        for item in requests:
            approved = self.request("post", f"/api/admin/import-requests/{item['id']}/approve", json={})
            self.assertEqual(approved.status_code, 200, approved.json)

        with self.app.app_context():
            db = self.app.get_db()
            assets = db.execute(
                """SELECT a.store_name, a.manager, a.product_line, a.equipment_type,
                          a.equipment_brand, a.model, a.area, a.province
                   FROM assets a ORDER BY a.id"""
            ).fetchall()
            self.assertEqual(len(assets), 2)
            self.assertEqual(assets[0]["store_name"], "EDITED FIRST STORE")
            self.assertEqual(assets[1]["store_name"], "NEW TEST STORE")
            self.assertEqual(assets[0]["manager"], "ROW MANAGER 1")
            self.assertEqual(assets[1]["manager"], "ROW MANAGER 2")
            self.assertEqual(assets[0]["product_line"], "ROW 1 PRODUCT_LINE")
            self.assertEqual(assets[1]["product_line"], "ROW 2 PRODUCT_LINE")
            self.assertEqual(assets[0]["area"], self.area)
            self.assertEqual(assets[0]["province"], self.province)
            self.assertEqual(assets[1]["area"], "")
            self.assertEqual(assets[1]["province"], "")


if __name__ == "__main__":
    unittest.main()
